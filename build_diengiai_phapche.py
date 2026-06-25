# -*- coding: utf-8 -*-
"""Diễn giải cho Pháp chế: các trường hợp GAPIT nhận DLCN của chủ thể từ Đại lý (gắn với DPA)."""
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT='/home/user/finalproject/output/Dien_giai_PhapChe_truong_hop_GAPIT_nhan_DLCN_tu_DaiLy.docx'
d=docx.Document()
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
RED=RGBColor(0xC0,0,0); NAVY=RGBColor(0x1A,0x2F,0x5A)

def H(t,sz=13,sb=10,center=False,color=NAVY):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=True; r.font.size=Pt(sz); r.font.color.rgb=color
    p.paragraph_format.space_before=Pt(sb); p.paragraph_format.space_after=Pt(4)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def P(t,just=True,it=False,bold=False,indent=None):
    p=d.add_paragraph(); r=p.add_run(t); r.italic=it; r.bold=bold
    if just: p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.left_indent=Cm(indent)
    p.paragraph_format.space_after=Pt(4); return p

H('DIỄN GIẢI CHO PHÁP CHẾ',sz=15,center=True,sb=0)
H('Các trường hợp GAPIT nhận Dữ liệu cá nhân (DLCN) của chủ thể dữ liệu TỪ Đại lý',sz=12,center=True,sb=2,color=RGBColor(0,0,0))
P('(Gắn với Thỏa thuận Xử lý DLCN – DPA của dịch vụ eSIM du lịch và dịch vụ zVAS; trả lời câu hỏi của Pháp chế về việc mô tả luồng dữ liệu trong mô hình đại lý)',just=False,it=True)

H('1. Nguyên tắc chung & vai trò các bên')
P('Trong mô hình đại lý, ĐẠI LÝ là bên có quan hệ trực tiếp với khách hàng/người dùng cuối và là bên thu thập DLCN ban đầu — do đó Đại lý là BÊN KIỂM SOÁT DỮ LIỆU (Bên A trong DPA).')
P('GAPIT chỉ tiếp nhận và xử lý DLCN của khách hàng/người dùng cuối ở mức hệ thống, theo chỉ dẫn của Đại lý, nhằm thực hiện nghĩa vụ cung cấp sản phẩm/dịch vụ — khi đó GAPIT là BÊN XỬ LÝ DỮ LIỆU (Bên B trong DPA).')
P('Quan trọng: KHÔNG phải mọi đơn hàng đều phát sinh việc GAPIT nhận DLCN của khách hàng cuối. DPA (kiểm soát–xử lý) chỉ thực sự được kích hoạt đối với các trường hợp CÓ chuyển giao DLCN nêu tại Mục 2. Ngoài các trường hợp đó, GAPIT chỉ xử lý DLCN của đầu mối/người liên hệ của Đại lý (quan hệ B2B), không phải DLCN của khách hàng cuối (Mục 3).',bold=False)

H('2. Các trường hợp GAPIT NHẬN DLCN của chủ thể từ Đại lý')
P('Đây là các trường hợp Đại lý chủ động chuyển/để GAPIT tiếp cận DLCN của khách hàng/người dùng cuối; DPA áp dụng cho các trường hợp này:',it=False)

rows=[
 ('TT','Trường hợp (khi nào)','Loại DLCN GAPIT nhận','Mục đích','Vai trò GAPIT'),
 ('A. DỊCH VỤ eSIM DU LỊCH','','','',''),
 ('a','Đại lý nhập email (và họ tên) của KHÁCH HÀNG CUỐI lên hệ thống CMS để GAPIT gửi trực tiếp mã QR/thông tin kích hoạt eSIM tới khách',
  'Email; họ tên; thông tin đơn hàng gắn với khách','Cấp và giao eSIM tới khách hàng cuối','Bên Xử lý (theo chỉ dẫn Đại lý)'),
 ('b','Khách hàng cuối yêu cầu xuất hóa đơn GTGT; Đại lý chuyển thông tin xuất hóa đơn của khách cho GAPIT',
  'Họ tên/tên tổ chức; mã số thuế cá nhân; địa chỉ; email nhận hóa đơn','Phát hành hóa đơn GTGT (qua VNPT) & lưu trữ kế toán',
  'Bên Xử lý; riêng mục đích hóa đơn–kế toán có thể là Bên Kiểm soát độc lập (nghĩa vụ luật thuế)'),
 ('c','Khiếu nại/hỗ trợ kỹ thuật vượt phạm vi Đại lý: Đại lý chuyển tiếp thông tin của khách để GAPIT/NCC xử lý',
  'Mã đơn; email; nội dung sự việc; ảnh chụp/bằng chứng','Chăm sóc khách hàng, khắc phục lỗi eSIM','Bên Xử lý'),
 ('B. DỊCH VỤ zVAS','','','',''),
 ('d','Đại lý nhờ GAPIT hỗ trợ kích hoạt/onboarding cho người dùng cuối; cung cấp thông tin định danh người dùng',
  'Họ tên; email; định danh tài khoản Zalo/ZingMP3 (để hỗ trợ kích hoạt)','Hỗ trợ kích hoạt Mã code & onboarding','Bên Xử lý'),
 ('đ','Người dùng cuối/Đại lý yêu cầu xuất hóa đơn GTGT (tương tự mục b)',
  'Họ tên/tên tổ chức; MST; địa chỉ; email nhận hóa đơn','Phát hành hóa đơn GTGT & lưu trữ kế toán','Như mục b'),
 ('e','Khiếu nại/hỗ trợ kỹ thuật vượt cấp; Đại lý chuyển tiếp thông tin để GAPIT xử lý/leo thang VNG',
  'Mã đơn/Mã code; email; nội dung sự việc','CSKH, xử lý sự cố kỹ thuật','Bên Xử lý'),
]
t=d.add_table(rows=len(rows),cols=5); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
widths=[Cm(0.9),Cm(6.2),Cm(4.5),Cm(3.8),Cm(3.6)]
for ri,row in enumerate(rows):
    cells=t.rows[ri].cells
    is_hdr=(ri==0); is_sec=row[1]=='' and row[0].startswith(('A.','B.'))
    for ci,val in enumerate(row):
        cell=cells[ci]; cell.width=widths[ci]
        cell.text=''
        pr=cell.paragraphs[0]; run=pr.add_run(val if not (is_sec and ci>0) else '')
        run.font.size=Pt(9.5 if not is_hdr else 9.5)
        if is_hdr: run.bold=True; run.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        if is_sec and ci==0: run.bold=True; run.font.color.rgb=NAVY
    if is_hdr:
        for c in cells:
            c._tc.get_or_add_tcPr().append(docx.oxml.parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="1B2F5A"/>'))
    if is_sec:
        # merge section row across columns
        a=cells[0]
        for c in cells[1:]:
            a=a.merge(c)
        a.paragraphs[0].runs[0].text=row[0] if a.paragraphs[0].runs else ''
        a._tc.get_or_add_tcPr().append(docx.oxml.parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="E6F1FB"/>'))

H('3. Các trường hợp GAPIT KHÔNG nhận DLCN của khách hàng cuối')
P('Trong các trường hợp sau, GAPIT KHÔNG tiếp nhận DLCN của khách hàng/người dùng cuối; Đại lý là bên kiểm soát duy nhất và tự chịu trách nhiệm đối với DLCN đó:')
for it in [
 'eSIM: Đại lý đặt mua theo lô, tự nhận mã QR/eSIM rồi tự phân phối và tự xuất hóa đơn cho khách của mình, KHÔNG nhập thông tin khách hàng cuối lên CMS của GAPIT.',
 'zVAS: Người dùng cuối TỰ kích hoạt Mã code tại zbox.vn trên tài khoản Zalo/ZingMP3 của chính họ; GAPIT không thu thập thông tin đăng nhập/tài khoản của người dùng cuối.',
 'Trong các trường hợp này, GAPIT chỉ xử lý DLCN của đầu mối/người liên hệ của Đại lý (B2B) để đặt hàng, thanh toán, xuất hóa đơn cho Đại lý.',
]:
    P('• '+it,indent=0.6)

H('4. Hệ quả pháp lý & khuyến nghị')
for it in [
 'DPA áp dụng theo mẫu PDPA(VNI): Đại lý = Bên Kiểm soát (Bên A), GAPIT = Bên Xử lý (Bên B). DPA điều chỉnh các trường hợp tại Mục 2.',
 'Đại lý phải bảo đảm đã có sự đồng ý hợp lệ của chủ thể dữ liệu (theo Luật BVDLCN 91/2025 & NĐ 356/2025) TRƯỚC khi chuyển DLCN cho GAPIT, và cung cấp bằng chứng đồng ý khi GAPIT yêu cầu.',
 'Riêng dữ liệu phục vụ hóa đơn/kế toán (mục b, đ): GAPIT lưu trữ theo pháp luật kế toán (tối thiểu 10 năm) và xử lý cho mục đích tuân thủ luật thuế — đối với mục đích này GAPIT có thể là Bên Kiểm soát độc lập; cần ghi rõ trong DPA/Phụ lục.',
 'Đề xuất: đưa chính xác bảng Mục 2 vào Phụ lục "Mô tả hoạt động xử lý dữ liệu" của DPA, kèm sơ đồ luồng dữ liệu (loại dữ liệu – bên thu thập – bên nhận – mục đích – nơi lưu trữ – bên thứ ba) — đúng nội dung Pháp chế yêu cầu mô tả.',
 'Nguyên tắc tối thiểu hóa: GAPIT chỉ yêu cầu/nhận từ Đại lý đúng các trường dữ liệu cần thiết cho từng mục đích nêu trên; không thu thập vượt phạm vi.',
]:
    P('• '+it,indent=0.6)

d.save(OUT); print('SAVED',OUT)
