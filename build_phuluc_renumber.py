# -*- coding: utf-8 -*-
"""Chỉnh hình thức Phụ lục Quy trình Mã code zVAS theo nhận xét:
- Đầu mục lớn -> Điều 1, Điều 2, Điều 3, Điều 4
- Mục con -> đánh số phân cấp 1.1, 2.1... ; điểm -> a) b) c) ; cấp sâu nhất -> (i) (ii)
- Bỏ ký tự bullet '•' và bỏ auto-numbering để dẫn chiếu điều khoản thuận tiện, có thụt lề phân cấp.
Giữ nguyên nội dung, bảng, chữ ký."""
import docx, os, re
from docx.oxml.ns import qn
from docx.shared import Inches

SRC = '/root/.claude/uploads/0ae05b9e-80f9-5907-8558-033a769d7c2e/d5ca5216-PhuLuc_QuyTrinh_MaCode_zVAS_VNG_GAPIT_revised_1.docx'
OUT = '/home/user/finalproject/output/PhuLuc_QuyTrinh_MaCode_zVAS_VNG_GAPIT_revised_2.docx'

d = docx.Document(SRC)
IND = {0: 0.0, 1: 0.30, 2: 0.62}

def set_text(p, txt):
    runs = p.runs
    if not runs:
        p.add_run(txt); return
    runs[0].text = txt
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def strip_num(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None:
        for np in pPr.findall(qn('w:numPr')):
            pPr.remove(np)

def clean(t):
    t = t.strip()
    if t.startswith('•'):
        t = t.lstrip('•').lstrip()
    t = re.sub(r'^Bước\s*\d+\s*:\s*', '', t)   # bỏ tiền tố "Bước N:" ở quy trình hoàn trả
    return t

# (match_substring, kind, arg, level, stripnum)
#   HEAD/SUB: arg = toàn văn mới (giữ bold sẵn có)
#   PT      : arg = nhãn ghép vào trước nội dung đã làm sạch
rules = [
 ('1. Mục đích và phạm vi',                 'HEAD', 'Điều 1. Mục đích và phạm vi', 0, False),
 ('Phụ lục này quy định trình tự',          'PT',   '1.1. ', 0, False),

 ('2. Sơ đồ tổng quát luồng Mã code',       'HEAD', 'Điều 2. Sơ đồ tổng quát luồng Mã code', 0, False),
 ('Bước 1: Nhận Mã code',                   'SUB',  '2.1. Bước 1 – Nhận Mã code (Bên B nhận từ Bên A)', 0, False),
 ('• Đặt hàng:',                            'PT',   'a) ', 1, False),
 ('• Thanh toán trả trước:',                'PT',   'b) ', 1, False),
 ('• Bàn giao Mã code:',                    'PT',   'c) ', 1, False),
 ('• Kiểm tra khi nhận:',                   'PT',   'd) ', 1, False),
 ('Bước 2: Xử lý và quản lý Mã code',       'SUB',  '2.2. Bước 2 – Xử lý và quản lý Mã code', 0, False),
 ('• Quản lý code pool:',                   'PT',   'a) ', 1, False),
 ('• Theo dõi tồn kho',                     'PT',   'b) ', 1, False),
 ('• Đối chiếu:',                           'PT',   'c) ', 1, False),
 ('• Trường hợp Khách hàng có khiếu nại',   'PT',   'd) ', 1, False),
 ('Ngoại trừ trường hợp Khách hàng vi phạm','PT',   'e) ', 1, False),
 ('Đối với trường hợp Mã code zVas gặp lỗi khiến', 'PT', '(i) ', 2, True),
 ('Đối với trường hợp Mã code Zvas đã được kích hoạt thành công', 'PT', '(ii) ', 2, True),
 ('Bước 3: Giao Mã code cho khách hàng',    'SUB',  '2.3. Bước 3 – Giao Mã code cho khách hàng (02 mô hình)', 0, False),
 ('Bước 4: Kích hoạt',                      'SUB',  '2.4. Bước 4 – Kích hoạt (Người dùng cuối)', 0, False),
 ('• Người dùng cuối kích hoạt',            'PT',   'a) ', 1, False),
 ('• Hỗ trợ:',                              'PT',   'b) ', 1, False),

 ('3. Xử lý Mã code lỗi và hỗ trợ',         'HEAD', 'Điều 3. Xử lý Mã code lỗi và hỗ trợ (dẫn chiếu Hợp đồng)', 0, False),
 ('• Đối với trường hợp Mã Code bị lỗi mà Bên B thông báo', 'PT', '3.1. ', 0, False),
 ('Mã Code đã được sử dụng:',               'PT',   'a) ', 1, True),
 ('Mã Code không thể kích hoạt mà Hai Bên', 'PT',   'b) ', 1, True),
 ('Mã Code kích hoạt không đúng dịch vụ',   'PT',   'c) ', 1, True),
 ('Trường hợp mã Code chưa kích hoạt mà đã hết hạn', 'PT', '3.2. ', 0, False),
 ('Đối với trường hợp Bên B chưa phân phối hết số lượng', 'PT', '3.3. ', 0, False),
 ('Bước 1: Bên B gửi thông báo bằng văn bản', 'PT', 'a) ', 1, False),
 ('Bước 2: Sau khi Bên A xác nhận chấp thuận', 'PT', 'b) ', 1, False),
 ('Bước 3: Hai Bên ký Biên bản xác nhận hoàn trả', 'PT', 'c) ', 1, False),

 ('4. Bảng tóm tắt quy trình',              'HEAD', 'Điều 4. Bảng tóm tắt quy trình', 0, False),
]

used = set()
for p in d.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    for idx, (m, kind, arg, lvl, sn) in enumerate(rules):
        if idx in used:
            continue
        if t.startswith(m) or m in t:
            if kind in ('HEAD', 'SUB'):
                set_text(p, arg)
            else:  # PT
                if sn:
                    strip_num(p)
                    p.style = d.styles['Normal']
                set_text(p, arg + clean(p.text))
            # cross-reference fix trong điểm d) của 2.1
            if 'Kiểm tra khi nhận' in m:
                set_text(p, p.text.replace('theo Mục 3 của Phụ lục', 'theo Điều 3 của Phụ lục'))
            p.paragraph_format.left_indent = Inches(IND[lvl])
            used.add(idx)
            break

os.makedirs('/home/user/finalproject/output', exist_ok=True)
d.save(OUT)
print('SAVED', OUT, '| matched', len(used), 'of', len(rules))
missing = [rules[i][0] for i in range(len(rules)) if i not in used]
print('MISSING:', missing)

# in lại toàn bộ để kiểm tra
d2 = docx.Document(OUT)
for p in d2.paragraphs:
    if p.text.strip():
        print('  ', p.text[:95])
