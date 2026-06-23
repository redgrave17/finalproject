# -*- coding: utf-8 -*-
"""Slide tổng hợp Mô hình & Quy trình kinh doanh zVAS (Roadmap & Training) cho GAPIT.
Lấy branding trực tiếp từ ClientviewGAPIT_zVAS_Portfolio.pptx (màu, font, logo GAPIT)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand palette (từ slide gốc) ----
ZALO   = RGBColor.from_string('0068FF')
NAVY   = RGBColor.from_string('001A40')
BLUE2  = RGBColor.from_string('003F88')
LBLUE  = RGBColor.from_string('00B4FF')
ORANGE = RGBColor.from_string('F5862C')
TEAL   = RGBColor.from_string('00C6AE')
GREEN  = RGBColor.from_string('22C55E')
PURPLE = RGBColor.from_string('7C3AED')
GRAY   = RGBColor.from_string('718096')
GRAY2  = RGBColor.from_string('4A5568')
LIGHT  = RGBColor.from_string('F7F8F8')
WHITE  = RGBColor.from_string('FFFFFF')
CARDLN = RGBColor.from_string('E2E8F0')
RISK   = RGBColor.from_string('E2574C')
HEAD='Arial'; BODY='Calibri'
LOGO='/home/user/finalproject/assets/image4.png'

prs=Presentation(); prs.slide_width=Inches(10); prs.slide_height=Inches(5.625)
s=prs.slides.add_slide(prs.slide_layouts[6])

def rect(x,y,w,h,fill,line=None,rounded=False,lw=0.75):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Inches(x),Inches(y),Inches(w),Inches(h))
    shp.shadow.inherit=False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    return shp

def txt(x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    for m in (tf.margin_left,): pass
    tf.margin_left=Inches(0.04); tf.margin_right=Inches(0.04); tf.margin_top=Inches(0.01); tf.margin_bottom=Inches(0.01)
    for i,(runs,sz,sp_before) in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(0); p.space_before=Pt(sp_before); p.line_spacing=1.0
        for t,c,b,fn in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b
            r.font.color.rgb=c; r.font.name=fn
    return tb

# ================= TOP BAR =================
rect(0,0,10,0.55,LIGHT)
s.shapes.add_picture(LOGO,Inches(0.45),Inches(0.10),height=Inches(0.36))
txt(1.62,0.10,2.4,0.36,[([('×  Zalo VAS',ZALO,True,HEAD)],13,0)],anchor=MSO_ANCHOR.MIDDLE)
rect(7.95,0.08,2.00,0.38,ORANGE,rounded=True)
txt(7.95,0.08,2.00,0.38,[([('ROADMAP & TRAINING',WHITE,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

# ================= TITLE =================
txt(0.50,0.60,9.0,0.40,[([('Mô Hình & Quy Trình Kinh Doanh Dịch Vụ zVAS',NAVY,True,HEAD)],22,0)])
txt(0.50,1.04,9.0,0.24,[([('Tổng hợp định hướng Business Roadmap & Đào tạo nội bộ GAPIT — hiểu dịch vụ → chuẩn hóa hợp đồng → quản trị thanh toán/rủi ro',GRAY,False,BODY)],9.5,0)])

# ================= ROW 1: 3 framing questions =================
cards=[
 (ZALO,'1','Đã có mô hình & quy trình kinh doanh dịch vụ này chưa?',
  '✓ CÓ — chuỗi phân phối & quy trình vận hành đã được chuẩn hóa (sơ đồ bên dưới).'),
 (BLUE2,'2','Cần mô hình để hiểu dịch vụ → mới xác định đúng hợp đồng?',
  '✓ ĐÚNG — mỗi quan hệ trong chuỗi gắn với MỘT loại hợp đồng riêng.'),
 (ORANGE,'3','Trả trước / trả sau chỉ là vấn đề thanh toán & rủi ro?',
  '✓ ĐÚNG — bản chất dịch vụ không đổi; chỉ khác cơ chế & rủi ro thanh toán.'),
]
cx=0.50; cw=3.00; gap=0.0
for i,(clr,num,q,a) in enumerate(cards):
    x=0.50+i*3.167
    rect(x,1.30,3.00,0.92,WHITE,line=CARDLN)
    rect(x,1.30,0.09,0.92,clr)
    txt(x+0.20,1.36,2.72,0.20,[([('?  ',clr,True,BODY),('Câu hỏi định hướng',GRAY,True,BODY)],8,0)])
    txt(x+0.20,1.55,2.72,0.40,[([(q,NAVY,True,BODY)],9.5,0)])
    txt(x+0.20,1.94,2.72,0.26,[([(a,clr,False,BODY)],8,0)])

# ================= SECTION 1: business model chain =================
txt(0.50,2.30,9.0,0.20,[([('1 · MÔ HÌNH KINH DOANH — CHUỖI PHÂN PHỐI & HỢP ĐỒNG TƯƠNG ỨNG',GRAY,True,BODY)],9.5,0)])
chain=[
 (BLUE2,'VNG / Zalo','Nhà phát hành zVAS'),
 (ZALO,'GAPIT','Đại lý ủy quyền (Reseller)'),
 (TEAL,'Đại lý cấp dưới / KHDN','Phân phối / Sử dụng'),
 (GREEN,'Người Dùng Cuối','Kích hoạt tại zbox.vn'),
]
bw=2.05; bx=0.50
for i,(clr,t1,t2) in enumerate(chain):
    x=0.50+i*2.36
    rect(x,2.54,bw,0.60,clr,rounded=True)
    txt(x,2.60,bw,0.30,[([(t1,WHITE,True,BODY)],10.5,0)],align=PP_ALIGN.CENTER)
    txt(x,2.88,bw,0.22,[([(t2,WHITE,False,BODY)],7.5,0)],align=PP_ALIGN.CENTER)
    if i<3:
        txt(x+bw,2.54,0.36,0.60,[([('→',clr,True,HEAD)],18,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
# contract captions under arrows
txt(0.50,3.16,9.0,0.22,[([('Hợp đồng tương ứng:  ',GRAY2,True,BODY),
   ('① HĐ Dịch vụ zVAS (VNG=A · GAPIT=B)   →   ② HĐ Đại lý zVAS / HĐ Cung cấp DV zVAS + DPA (Khách hàng/Đại lý=A · GAPIT=B)',NAVY,False,BODY)],8.5,0)])

# ================= SECTION 2 (left): operating process =================
rect(0.50,3.52,4.40,1.62,WHITE,line=CARDLN)
rect(0.50,3.52,0.09,1.62,ZALO)
txt(0.70,3.58,4.10,0.22,[([('2 · QUY TRÌNH VẬN HÀNH (theo từng đơn)',ZALO,True,BODY)],9.5,0)])
steps=[
 ('1','Đặt hàng trên CMS / email được chỉ định'),
 ('2','Thanh toán — TRẢ TRƯỚC 100% theo từng đơn'),
 ('3','Bàn giao Mã code (≤ 03 ngày làm việc) → kết thúc đơn'),
 ('4','Người Dùng Cuối kích hoạt tại zbox.vn/activate-code'),
 ('5','CSKH: GAPIT xử lý cấp 1 → leo thang VNG cấp 2'),
]
for i,(n,t) in enumerate(steps):
    yy=3.82+i*0.265
    rect(0.72,yy,0.22,0.20,ZALO,rounded=True)
    txt(0.72,yy-0.005,0.22,0.21,[([(n,WHITE,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(1.00,yy-0.01,3.85,0.22,[([(t,GRAY2,False,BODY)],8.7,0)],anchor=MSO_ANCHOR.MIDDLE)

# ================= SECTION 3 (right): payment & risk =================
rect(5.00,3.52,4.50,1.62,WHITE,line=CARDLN)
rect(5.00,3.52,0.09,1.62,ORANGE)
txt(5.20,3.58,4.20,0.22,[([('3 · THANH TOÁN & RỦI RO',ORANGE,True,BODY)],9.5,0)])
# two mini cards
rect(5.22,3.84,2.05,0.78,RGBColor.from_string('EAF7EF'),line=GREEN,rounded=True)
txt(5.34,3.88,1.85,0.20,[([('TRẢ TRƯỚC (mặc định)',GREEN,True,BODY)],8.5,0)])
txt(5.34,4.07,1.85,0.52,[([('Bàn giao code SAU khi thanh toán đủ → kết thúc đơn; không công nợ → rủi ro thấp.',GRAY2,False,BODY)],7.6,0)])
rect(7.36,3.84,2.05,0.78,RGBColor.from_string('FDEEEC'),line=RISK,rounded=True)
txt(7.48,3.88,1.85,0.20,[([('TRẢ SAU (công nợ)',RISK,True,BODY)],8.5,0)])
txt(7.48,4.07,1.85,0.52,[([('Phát sinh công nợ → rủi ro thu hồi; cần hạn mức/bảo lãnh & lãi chậm trả (nếu thỏa thuận).',GRAY2,False,BODY)],7.6,0)])
# principle line
rect(5.22,4.70,4.19,0.40,RGBColor.from_string('FFF6E9'),line=ORANGE,rounded=True)
txt(5.34,4.71,3.95,0.38,[([('★ ',ORANGE,True,BODY),('Trả trước/sau KHÔNG đổi bản chất dịch vụ — chỉ khác cơ chế & rủi ro thanh toán. Mã code đã mua: không hoàn/đổi (chính sách VNG).',NAVY,False,BODY)],7.8,0)],anchor=MSO_ANCHOR.MIDDLE)

# ================= FOOTER =================
rect(0.50,5.30,9.00,0.22,BLUE2,rounded=True)
txt(0.65,5.30,8.70,0.22,[([('GAPIT JSC  |  gapit.com.vn',WHITE,True,BODY),
   ('     —  Tài liệu định hướng Business Roadmap & Đào tạo nội bộ',WHITE,False,BODY)],8,0)],anchor=MSO_ANCHOR.MIDDLE)

import os
os.makedirs('/home/user/finalproject/output',exist_ok=True)
out='/home/user/finalproject/output/GAPIT_zVAS_MoHinh_QuyTrinh_Roadmap.pptx'
prs.save(out); print('SAVED',out)
