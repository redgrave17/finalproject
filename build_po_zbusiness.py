# -*- coding: utf-8 -*-
"""Điều chỉnh PO mẫu (vốn cho dịch vụ ĐỊNH DANH/SMS) thành PO cho khách hàng mua 1 gói zBusiness.
- Bỏ nội dung đặc thù 'định danh', thay bằng dịch vụ/bàn giao Mã code gói zBusiness.
- VAT 10%; đơn giá 1.990.000 (chưa VAT) -> Tổng tiền 2.189.000. KHÔNG ghi chú ngoại lệ giá chưa VAT.
- Giữ nguyên định dạng (chỉ đổi .value)."""
import openpyxl, os

SRC = '/root/.claude/uploads/0ae05b9e-80f9-5907-8558-033a769d7c2e/7d3e4f4d-PO.xlsx'
OUT = '/home/user/finalproject/output/PO_zBusiness.xlsx'

wb = openpyxl.load_workbook(SRC)
ws = wb['PO_DINH DANH']

# --- PO number: đổi hậu tố dịch vụ .sms -> .zVAS ---
ws['D3'] = 'PO No: 041/2026/SỮA VIỆT NAM-GAPIT/S.zVAS'

# --- Mục I: tiêu đề & hạng mục ---
ws['A13'] = 'I. CHI PHÍ DỊCH VỤ GÓI zBUSINESS'
ws['A15'] = 'HẠNG MỤC DỊCH VỤ zBUSINESS'
ws['B16'] = 'Gói zBusiness (12 tháng)'
ws['E16'] = 'Gói'

# --- Giá & VAT 10% ---
# THÀNH TIỀN = ĐƠN GIÁ x SỐ LƯỢNG (chưa VAT); TỔNG TIỀN = có VAT 10%
ws['C16'] = 1990000
ws['F16'] = '=C16*D16'
ws['F17'] = '=SUM(F16:F16)*1.1'   # 1.990.000 x 1.1 = 2.189.000

# --- Ghi chú: bỏ dòng "VAT 8%"; chỉ nêu VAT 10% trên tổng + thời hạn gói ---
ws['A18'] = ('GHI CHÚ:\n'
             '1. Tổng tiền đã bao gồm thuế GTGT (VAT) 10%.\n'
             '2. Gói zBusiness có thời hạn sử dụng 12 (mười hai) tháng kể từ ngày kích hoạt Mã code.')

# --- Mục II: thay nội dung đặc thù định danh bằng bàn giao Mã code zBusiness ---
ws['A22'] = ('2.1 Thời gian bắt đầu cung cấp dịch vụ gói zBusiness kể từ ngày đơn đặt hàng này '
             'được ký kết.')
ws['A28'] = ('Bên A có trách nhiệm cung cấp và bàn giao cho Bên B: Mã code kích hoạt gói dịch vụ '
             'zBusiness, gửi qua địa chỉ email Bên B đã đăng ký. Người Dùng Cuối kích hoạt Mã code '
             'theo hướng dẫn của Bên A.')
ws['A29'] = ('2.3 Điều kiện thanh toán:\n'
             'Bên B tiến hành thanh toán 100% giá trị phí dịch vụ gói zBusiness cho GAPIT trong vòng '
             '45 ngày đến 90 ngày kể từ ngày đơn đặt hàng này được ký kết và hồ sơ thanh toán hợp lệ '
             'bao gồm:\n'
             'a) Biên bản bàn giao nghiệm thu\n'
             'b) Hóa đơn Giá trị gia tăng hợp lệ tương ứng 100% phí dịch vụ\n')

os.makedirs('/home/user/finalproject/output', exist_ok=True)
wb.save(OUT)
print('SAVED', OUT)

# verify
wb2 = openpyxl.load_workbook(OUT)
w2 = wb2['PO_DINH DANH']
for coord in ['D3','A13','A15','B16','C16','F16','A17','F17','A18','A22','A28','A29']:
    print(coord, '=>', repr(w2[coord].value)[:90])
