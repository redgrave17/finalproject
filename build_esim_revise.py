# -*- coding: utf-8 -*-
"""Sửa HĐ Đại lý eSIM theo góp ý pháp chế (giữ GAPIT=Bên A). Highlight phần sửa."""
import copy, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor, Pt
Y=WD_COLOR_INDEX.YELLOW; RED=RGBColor(0xC0,0,0)
SRC='/root/.claude/uploads/0ae05b9e-80f9-5907-8558-033a769d7c2e/9f5b623d-HopDong_DaiLyeSIM.docx'
OUT='/home/user/finalproject/output/HopDong_DaiLy_eSIM_SUA_THEO_GOPY.docx'
d=docx.Document(SRC)
plist=list(d.paragraphs)
def find(needle):
    for p in plist:
        if needle in p.text: return p
    raise ValueError('NOT FOUND: '+needle)
def setp(needle,new):
    p=find(needle)
    if p.runs:
        p.runs[0].text=new
        for r in p.runs[1:]: r.text=''
        r0=p.runs[0]
    else: r0=p.add_run(new)
    r0.font.highlight_color=Y
def insert_after(needle,text,note=False):
    p=find(needle)
    np_el=OxmlElement('w:p'); p._p.addnext(np_el)
    pPr=p._p.find(qn('w:pPr'))
    if pPr is not None:
        new_pPr=copy.deepcopy(pPr)
        for nm in new_pPr.findall(qn('w:numPr')): new_pPr.remove(nm)
        np_el.append(new_pPr)
    np=Paragraph(np_el,p._parent)
    if note:
        r=np.add_run('【LƯU Ý PHÁP CHẾ】 '); r.bold=True; r.font.color.rgb=RED; r.font.size=Pt(10)
    r2=np.add_run(text); r2.font.highlight_color=Y; r2.font.size=Pt(13); r2.font.name='Times New Roman'
    return np

# 1) GAPIT là bên giao đại lý — bổ sung quyền bổ nhiệm đại lý + lưu ý rà soát HĐ Consortio
setp('sở hữu và vận hành các kênh kinh doanh trực tuyến',
 'Bên A là doanh nghiệp được thành lập hợp pháp tại Việt Nam, sở hữu và vận hành các kênh kinh doanh trực tuyến '
 '(bao gồm Zalo Mini App), có quyền phân phối và được phép phát triển, bổ nhiệm đại lý để phân phối các sản phẩm '
 'SIM/eSIM du lịch tới khách hàng;')
insert_after('có quyền phân phối và được phép phát triển, bổ nhiệm đại lý',
 'Đề nghị rà soát Hợp đồng giữa GAPIT và nhà cung cấp eSIM (Consortio) để xác nhận GAPIT có quyền phân phối lại '
 'và bổ nhiệm đại lý đối với Sản phẩm. Nếu hợp đồng với Consortio chưa quy định rõ quyền này, cần bổ sung/lấy xác '
 'nhận bằng văn bản từ Consortio trước khi ký Hợp đồng đại lý này nhằm bảo đảm tính hợp lệ của quan hệ giao đại lý.',
 note=True)

# 2) 2.3 — bổ sung đường link Zalo Mini App
setp('“Giá bán lẻ niêm yết trên Zalo Mini App” (“Giá Zalo Mini App”) là giá bán lẻ công khai',
 '2.3. “Giá bán lẻ niêm yết trên Zalo Mini App” (“Giá Zalo Mini App”) là giá bán lẻ công khai của Sản phẩm do '
 'Bên A niêm yết trên kênh Zalo Mini App chính thức của Bên A (link: https://zalo.me/s/4552695804193899018/) '
 'tại từng thời điểm.')

# 3) 4.1 & 4.3 — làm rõ, hết mâu thuẫn (15 ngày để thống nhất; 30 ngày tự động nếu không thỏa thuận lại)
setp('cam kết không thay đổi Giá cấp cho Đại lý trong tối thiểu 90',
 'Bên A cam kết không thay đổi Giá cấp cho Đại lý trong tối thiểu 90 (chín mươi) ngày kể từ ngày áp dụng. Trường '
 'hợp điều chỉnh, Bên A thông báo cho Bên B tối thiểu 15 (mười lăm) ngày làm việc trước khi áp dụng để Hai Bên '
 'thống nhất lại giá; giá bán cuối cùng sau điều chỉnh và thỏa thuận được Bên A thông báo và ghi nhận qua email/CMS.')
setp('Bên A có quyền cập nhật Giá cấp cho Đại lý, Giá bán khuyến nghị',
 'Bên A có quyền cập nhật Giá cấp cho Đại lý, Giá bán khuyến nghị và Giá niêm yết công khai trên Zalo Mini App '
 'theo từng thời điểm và thông báo cho Bên B qua email/CMS. Trong thời gian thông báo (tối thiểu 15 ngày làm việc '
 'theo Điều 4.1), Hai Bên có thể thương lượng lại giá; nếu Hai Bên không đạt được thỏa thuận khác, mức giá mới tự '
 'động có hiệu lực sau 30 (ba mươi) ngày kể từ ngày Bên A thông báo.')

# 4) 5.4 — theo đúng văn bản chị đưa
setp('Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ tương ứng với từng đơn hàng đã thanh toán',
 '5.4. Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ tương ứng với từng đơn hàng và kể từ ngày tiền thanh toán được '
 'ghi có vào tài khoản của Bên A mở tại Ngân hàng (theo thông tin tài khoản thanh toán quy định trong Hợp đồng) '
 'theo quy định của pháp luật.')

# 5) 5.5 — theo đúng văn bản chị đưa
setp('Tra cứu và đối chiếu: Bên B chủ động tra cứu lịch sử đơn hàng',
 '5.5. Tra cứu và đối chiếu: Bên B chủ động tra cứu lịch sử đơn hàng, trạng thái xử lý và hóa đơn trên CMS. Hai '
 'Bên không tiến hành đối soát định kỳ; trường hợp có sai lệch, Hai Bên rà soát và điều chỉnh trên tinh thần thiện chí.')

# 6) đánh số lại 5.6 -> 5.7 (Hủy đơn) để chèn điều khoản bàn giao thành 5.6
setp('Hủy đơn và hoàn tiền: Sau khi eSIM đã được cấp mã QR',
 '5.7. Hủy đơn và hoàn tiền: Sau khi eSIM đã được cấp mã QR hoặc thông tin kích hoạt, đơn hàng không được hủy hoặc '
 'hoàn tiền, trừ trường hợp lỗi được xác định thuộc trách nhiệm của Bên A.')

# 7) chèn điều khoản BÀN GIAO eSIM (khi nào + cách thức) sau 5.5
insert_after('Hai Bên không tiến hành đối soát định kỳ; trường hợp có sai lệch',
 '5.6. Bàn giao Sản phẩm (eSIM): Bên A chỉ cấp và bàn giao mã QR/thông tin kích hoạt eSIM SAU KHI tiền thanh toán '
 'của đơn hàng được ghi có đầy đủ vào tài khoản của Bên A. Hình thức bàn giao: qua hệ thống CMS và/hoặc gửi tới '
 'email do Bên B đăng ký (hoặc gửi tới email khách hàng cuối khi Bên B chỉ định), trong thời gian sớm nhất kể từ '
 'khi xác nhận thanh toán (chi tiết quy trình và thời gian tại Phụ lục 02).')

# 8) Phụ lục 01 — ghi rõ giá đã/chưa bao gồm GTGT
insert_after('Giá cấp cho Đại lý và chiết khấu là điều kiện thương mại nội bộ',
 'Ghi chú thuế: Toàn bộ đơn giá tại Bảng giá Sản phẩm là giá ĐÃ BAO GỒM thuế giá trị gia tăng (GTGT) 10%. '
 'Trường hợp áp dụng đơn giá chưa bao gồm GTGT, phải ghi rõ tại từng dòng tương ứng.')

# 9) 8.3 — viết lại gọn (bỏ ràng buộc số lượng đặt hàng tối thiểu) theo góp ý
setp('Nguồn cung ứng và ràng buộc số lượng đơn đặt hàng',
 '8.3. Nguồn cung ứng: Bên A là nhà cung cấp chính thức của Bên B đối với Sản phẩm theo Hợp đồng này.')

# 10) Phụ lục 01 — note chính sách chiết khấu theo năm/doanh thu (thay cho ràng buộc tại 8.3)
insert_after('Áp dụng cho tất cả sản phẩm/gói cước',
 'Chính sách chiết khấu nêu trên áp dụng cho năm đầu tiên của Hợp đồng. Từ năm thứ hai trở đi, nếu Bên B không đạt '
 'doanh thu tối thiểu …… VNĐ/năm (hoặc mức sản lượng do Hai Bên thỏa thuận), Bên A có quyền xem xét, điều chỉnh lại '
 'tỷ lệ chiết khấu áp dụng cho Bên B.')

# 11) Xóa đoạn 8.3 cũ (cam kết không cạnh tranh, đã gạch ngang) cho gọn
for _p in list(d.paragraphs):
    if 'kể từ ngày chấm' in _p.text and 'KHÔNG, trực tiếp hoặc gián tiếp' in _p.text:
        _p._p.getparent().remove(_p._p); break


d.save(OUT); print('SAVED',OUT)
