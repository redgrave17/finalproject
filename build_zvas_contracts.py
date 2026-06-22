# -*- coding: utf-8 -*-
"""Điều chỉnh 2 hợp đồng zVAS theo nhận xét, highlight phần thay đổi.
- Đảo vai trò: Khách hàng/Đại lý = Bên A, GAPIT = Bên B (reorder khối chủ thể).
- Áp các sửa đổi theo từng nhận xét.
Cơ chế: sửa điều khoản theo quy ước HIỆN TẠI (GAPIT=Bên A) rồi swap A<->B toàn văn
(trừ khối chủ thể đã viết theo quy ước MỚI và các ghi chú chèn sau)."""
import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

UP = '/root/.claude/uploads/0ae05b9e-80f9-5907-8558-033a769d7c2e/'
Y = WD_COLOR_INDEX.YELLOW

def first_fmt(p):
    for r in p.runs:
        if r.text.strip():
            return r
    return p.runs[0] if p.runs else None

def set_text_hl(p, text, hl=True):
    fmt = first_fmt(p)
    bold = fmt.bold if fmt else None
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r.text = ''
        r0 = p.runs[0]
    else:
        r0 = p.add_run(text)
    r0.bold = bold
    if hl: r0.font.highlight_color = Y

def swap_ab(s):
    s = s.replace('Bên A', '\x00').replace('Bên B', 'Bên A').replace('\x00', 'Bên B')
    s = s.replace('BÊN A', '\x01').replace('BÊN B', 'BÊN A').replace('\x01', 'BÊN B')
    return s

def swap_para(p):
    full = p.text
    sw = swap_ab(full)
    if sw == full:
        return
    fmt = first_fmt(p)
    bold = fmt.bold if fmt else None
    size = fmt.font.size if fmt else None
    name = fmt.font.name if fmt else None
    hl = fmt.font.highlight_color if fmt else None
    p.runs[0].text = sw
    for r in p.runs[1:]: r.text = ''
    r0 = p.runs[0]
    r0.bold = bold
    if size: r0.font.size = size
    if name: r0.font.name = name
    if hl: r0.font.highlight_color = hl

def insert_note_after(p, text):
    new_p = OxmlElement('w:p')
    p._p.addnext(new_p)
    np = Paragraph(new_p, p._parent)
    r = np.add_run('【ĐỀ XUẤT/GÓP Ý】 ')
    r.bold = True; r.font.color.rgb = RGBColor(0xC0,0,0); r.font.size = Pt(10)
    r2 = np.add_run(text); r2.font.highlight_color = Y
    return np

def process(src, out, edits, party_start_needle, party_lines, notes):
    d = docx.Document(src)
    plist = list(d.paragraphs)
    def fp(needle):
        for p in plist:
            if needle in p.text: return p
        raise ValueError('NOT FOUND: '+needle)

    # 1) Sửa điều khoản theo quy ước hiện tại (GAPIT=Bên A) + highlight
    for needle, newtext in edits:
        set_text_hl(fp(needle), newtext, hl=True)

    # 2) Reorder + viết lại khối chủ thể theo quy ước MỚI (Bên A=khách hàng/đại lý)
    pstart = fp(party_start_needle)
    sidx = plist.index(pstart)
    excluded = set()
    for k, line in enumerate(party_lines):
        p = plist[sidx + k]
        set_text_hl(p, line, hl=True)
        excluded.add(id(p._p))

    # 3) Swap A<->B toàn văn, trừ khối chủ thể
    for p in d.paragraphs:
        if id(p._p) in excluded: continue
        swap_para(p)
    for t in d.tables:
        for row in t.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                for p in c.paragraphs:
                    swap_para(p)

    # 4) Chèn ghi chú (sau swap, theo quy ước mới)
    plist2 = list(d.paragraphs)
    def fp2(needle):
        for p in plist2:
            if needle in p.text: return p
        raise ValueError('NOTE ANCHOR NOT FOUND: '+needle)
    for anchor, text in notes:
        insert_note_after(fp2(anchor), text)

    d.save(out)
    print('SAVED', out)

# =========================================================================
# F1 — HỢP ĐỒNG ĐẠI LÝ zVAS (GAPIT v Đại lý)
# =========================================================================
f1_edits = [
 # Comment 3: Điều 4.1 — 15 ngày làm việc -> 30 ngày (đồng bộ Điều 4.3)
 ('Bên A cam kết không thay đổi Giá cấp cho Đại lý và mức chiết khấu trong tối thiểu 90',
  'Bên A cam kết không thay đổi Giá cấp cho Đại lý và mức chiết khấu trong tối thiểu 90 (chín mươi) ngày kể từ '
  'ngày áp dụng. Trường hợp điều chỉnh, Bên A thông báo cho Bên B tối thiểu 30 (ba mươi) ngày trước khi giá mới '
  'có hiệu lực (đồng bộ với thời hạn hiệu lực điều chỉnh giá tại Điều 4.3).'),
 # Comment 2: Điều 4.2 — chiết khấu không được phá giá sàn Zalo
 ('Chiết khấu từ biên lợi nhuận của chính Đại lý',
  'Ưu đãi từ biên lợi nhuận của chính Đại lý (không công khai): Ràng buộc nêu trên áp dụng đối với giá công bố/'
  'niêm yết công khai. Bên B có thể dành ưu đãi cho khách hàng từ phần chiết khấu/biên lợi nhuận của chính mình '
  'trên cơ sở thương lượng riêng (không công khai), với điều kiện trong mọi trường hợp KHÔNG niêm yết, quảng cáo, '
  'chào bán hoặc bán Mã code zVAS ở mức giá thấp hơn Giá bán lẻ trực tiếp của Zalo (Giá sàn công khai) trên bất kỳ '
  'kênh nào, nhằm tuân thủ quy định của nhà phát hành (VNG). Mọi hình thức giảm giá khiến giá bán thực tế thấp hơn '
  'Giá sàn công khai chỉ được thực hiện khi có chấp thuận trước bằng văn bản của Bên A.'),
 # Comment 5: Điều 5.5 — thời điểm xuất hóa đơn tính từ ngày tiền ghi có
 ('Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ (thuế suất 10%)',
  '5.5. Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ (thuế suất 10%) tương ứng với từng đơn hàng đã thanh toán '
  'theo quy định của pháp luật, trong vòng 03–05 ngày làm việc kể từ ngày tiền thanh toán được ghi có vào tài khoản '
  'của Bên A mở tại Ngân hàng theo thông tin tài khoản quy định trong Hợp đồng.'),
 # Comment 4: Điều 5.6 — trả trước, kết thúc đơn hàng, không đối soát hàng tháng
 ('Tra cứu và đối chiếu: Bên B chủ động tra cứu',
  '5.6. Bản chất trả trước và kết thúc đơn hàng: Hợp đồng này vận hành theo cơ chế trả trước theo từng đơn; sau khi '
  'Bên B thanh toán đủ và Bên A bàn giao Mã code zVAS theo Điều 5.3, đơn hàng được coi là hoàn tất (kết thúc đơn '
  'hàng). Bên B chủ động tra cứu lịch sử đơn hàng, trạng thái xử lý và hóa đơn trên CMS/zbox.vn; Hai Bên KHÔNG thực '
  'hiện đối soát hàng tháng. Trường hợp có sai lệch, Hai Bên rà soát và điều chỉnh trên tinh thần thiện chí.'),
 # Comment 6a: Điều 9.3 — không hoàn tiền mã code chưa kích hoạt
 ('Hoàn tiền Mã code chưa kích hoạt: Trường hợp Bên B không có nhu cầu',
  '9.3. Mã code đã thanh toán không hoàn tiền: Phù hợp với chính sách của nhà phát hành (VNG), Mã code zVAS đã '
  'thanh toán sẽ không được hoàn tiền dưới bất kỳ hình thức nào, kể cả đối với Mã code chưa kích hoạt; ngoại trừ '
  'trường hợp Mã code lỗi thuộc trách nhiệm của Bên A được xử lý theo Điều 9.1. (Lý do: Bên A không được nhà phát '
  'hành cho phép hoàn/đổi Mã code đã thanh toán nên không nhận ràng buộc nghĩa vụ này đối với Bên B.)'),
 # Comment 6b: Điều 12.3 — bỏ hoàn trả mã code chưa kích hoạt khi chấm dứt
 ('Hoàn trả Mã code chưa kích hoạt: Đối với các Mã code zVAS đã được Bên A bàn giao',
  '12.3. Xử lý Mã code đã bàn giao khi chấm dứt: Đối với các Mã code zVAS đã được Bên A bàn giao và Bên B đã thanh '
  'toán, khi chấm dứt Hợp đồng sẽ không được hoàn trả hoặc hoàn tiền, phù hợp với chính sách không hoàn/đổi của nhà '
  'phát hành (VNG) và Điều 9.3. Bên B chủ động sử dụng hoặc phân phối hết các Mã code đã nhận trước thời điểm chấm dứt.'),
]
# Khối chủ thể F1: paras bắt đầu "1. BÊN GIAO ĐẠI LÝ (BÊN A): CÔNG TY CỔ PHẦN GAPIT" -> reorder Đại lý=Bên A trước
f1_party_start = '1. BÊN GIAO ĐẠI LÝ (BÊN A): CÔNG TY CỔ PHẦN GAPIT'
f1_party = [
 '1. ĐẠI LÝ (BÊN A): ……………………………………………………………',
 'Trụ sở: ……………………………………………………………………………………',
 'Điện thoại: ……………………………      Mã số thuế: ……………………………',
 'Đại diện theo pháp luật: ……………………………      Chức vụ: ……………………',
 '(sau đây gọi tắt là “Bên A” hoặc “Đại lý”); và',
 '2. BÊN GIAO ĐẠI LÝ (BÊN B): CÔNG TY CỔ PHẦN GAPIT',
 'Trụ sở: Phòng 902, tầng 9, D10 Giảng Võ, Phường Giảng Võ, Thành phố Hà Nội, Việt Nam',
 'Điện thoại: 02435121928',
 'Mã số thuế: 0101816595',
 'Đại diện theo pháp luật: Ông Nguyễn Văn Long',
 'Chức vụ: Giám đốc',
 '(sau đây gọi tắt là “Bên B” hoặc “Bên giao đại lý”)',
 'Bên A và Bên B sau đây gọi riêng là “Bên”, gọi chung là “Hai Bên”.',
]
f1_notes = [
 ('Authorized Reseller Agent',
  'Cần xác nhận: đề nghị đính kèm/bổ sung văn bản chứng nhận GAPIT là Đại lý phân phối được ủy quyền chính thức '
  '(Authorized Reseller Agent) của VNG đối với Dịch vụ zVAS (Zalo, ZingMP3) làm cơ sở pháp lý cho tư cách phân phối.'),
 ('Bản chất trả trước và kết thúc đơn hàng',
  'Khuyến nghị áp dụng thống nhất các điều khoản trả trước theo mẫu “mua hàng trả trước” do Công ty đã ban hành '
  'để bảo đảm nhất quán giữa các hợp đồng.'),
]

# =========================================================================
# F2 — HỢP ĐỒNG CUNG CẤP DỊCH VỤ zVAS (Khách hàng DN)
# =========================================================================
f2_edits = [
 # [7] Điều 5.5 — trả trước, không lãi chậm trả trừ khi thỏa thuận
 ('Chậm thanh toán: Trường hợp Bên B chậm thanh toán quá 05',
  '5.5. Cơ chế trả trước: Hợp đồng vận hành theo cơ chế trả trước — Bên A chỉ bàn giao/kích hoạt Mã code sau khi '
  'Bên B đã thanh toán đủ theo Điều 5.3. Do đó Hai Bên không áp dụng lãi chậm thanh toán; trường hợp Bên B chưa '
  'thanh toán thì Bên A chưa có nghĩa vụ bàn giao/kích hoạt Mã code. Hai Bên chỉ áp dụng lãi chậm thanh toán nếu '
  'có thỏa thuận riêng bằng văn bản.'),
 # [12] Điều 10.1 — không hoàn tiền mã code chưa kích hoạt
 ('Mã code chưa kích hoạt: Trường hợp Bên B không có nhu cầu',
  '10.1. Mã code chưa kích hoạt: Phù hợp với chính sách của nhà phát hành (VNG), Mã code zVAS đã thanh toán không '
  'được hoàn tiền dưới bất kỳ hình thức nào, kể cả đối với Mã code chưa kích hoạt. Trường hợp Mã code lỗi thuộc '
  'trách nhiệm của Bên A được xử lý theo Điều 7.2.'),
 # [21] Điều 16.4 — bỏ hoàn trả mã code chưa kích hoạt khi chấm dứt
 ('Đối với Mã code đã bàn giao nhưng chưa kích hoạt tại thời điểm chấm dứt: Bên B được trả lại',
  '16.4. Khi chấm dứt Hợp đồng, Hai Bên hoàn tất các nghĩa vụ tài chính còn tồn đọng, xử lý khiếu nại (nếu có) và '
  'thực hiện thanh lý theo Điều 17. Theo chính sách chung, Mã code zVAS đã mua (đã thanh toán) không được hoàn trả '
  'hoặc hoàn tiền dưới mọi hình thức, kể cả Mã code đã bàn giao nhưng chưa kích hoạt; Bên B chủ động sử dụng hết Mã '
  'code đã nhận trước thời điểm chấm dứt.'),
 # [14] Điều 11.2 — điền nội dung bảo vệ dữ liệu cá nhân (đang trống)
 ('11.2. Bảo vệ dữ liệu cá nhân:',
  '11.2. Bảo vệ dữ liệu cá nhân: Trường hợp việc thực hiện Hợp đồng phát sinh hoạt động xử lý dữ liệu cá nhân (ví dụ '
  'thông tin của Người Dùng Cuối do Bên B chỉ định phục vụ kích hoạt, hỗ trợ), Hai Bên xác định rõ vai trò Bên Kiểm '
  'soát/Bên Xử lý và ký kết Thỏa thuận Xử lý Dữ liệu Cá nhân (DPA) là một phần không tách rời của Hợp đồng. Mỗi Bên '
  'tuân thủ Luật Bảo vệ Dữ liệu Cá nhân số 91/2025/QH15, Nghị định số 356/2025/NĐ-CP và pháp luật có liên quan; áp '
  'dụng biện pháp kỹ thuật, tổ chức phù hợp để bảo vệ dữ liệu cá nhân, chỉ xử lý đúng mục đích và bảo đảm quyền của '
  'chủ thể dữ liệu (truy cập, chỉnh sửa, xóa, rút lại sự đồng ý).'),
]
f2_party_start = 'BÊN CUNG CẤP DỊCH VỤ (BÊN A): CÔNG TY CỔ PHẦN GAPIT'
f2_party = [
 'BÊN SỬ DỤNG DỊCH VỤ (BÊN A): ……………………………………………………',
 'Trụ sở: ……………………………………………………………………………………',
 'Điện thoại: ……………………………      Mã số thuế: ……………………………',
 'Số tài khoản: ……………………………      Tại ngân hàng: ……………………………',
 'Đại diện theo pháp luật: ……………………………      Chức vụ: ……………………',
 'Giấy ủy quyền (nếu có): ……………………………………………………………………',
 '(sau đây gọi tắt là “Bên A” hoặc “Khách hàng”); và',
 'BÊN CUNG CẤP DỊCH VỤ (BÊN B): CÔNG TY CỔ PHẦN GAPIT',
 'Trụ sở: Phòng 902, tầng 9, D10 Giảng Võ, Phường Giảng Võ, Thành phố Hà Nội, Việt Nam',
 'Địa chỉ giao dịch: Tầng 21, tòa Vinaconex – Số 459C Bạch Mai, phường Bạch Mai, Thành phố Hà Nội',
 'Điện thoại: 02435121928 – Hotline: 1900 633 155',
 'Mã số thuế: 0101816595',
 'Số tài khoản: 110601258000 – Ngân hàng TMCP Công Thương Việt Nam (Vietinbank)',
 'Đại diện theo pháp luật: Ông Nguyễn Văn Long – Chức vụ: Giám đốc',
 '(sau đây gọi tắt là “Bên B” hoặc “Bên cung cấp”)',
 'Bên A và Bên B sau đây gọi riêng là “Bên”, gọi chung là “Hai Bên”.',
]
f2_notes = [
 ('Authorized Reseller Agent',
  'Cần xác nhận: đề nghị đính kèm/bổ sung văn bản chứng nhận GAPIT là Đại lý phân phối được ủy quyền chính thức '
  '(Authorized Reseller Agent) của VNG đối với Dịch vụ zVAS làm cơ sở pháp lý cho tư cách cung cấp dịch vụ.'),
 ('tranh chấp được đưa ra giải quyết tại Tòa án nhân dân có thẩm quyền',
  'Đã chọn Tòa án nhân dân có thẩm quyền (không dùng trọng tài) theo góp ý nhằm tránh án phí cao và bất lợi khi '
  'thi hành án.'),
]

import os
os.makedirs('/home/user/finalproject/output', exist_ok=True)
process(UP+'7a827907-HopDong_DaiLy_zVAS_GAPITvDaiLy_1.docx',
        '/home/user/finalproject/output/HopDong_DaiLy_zVAS_GAPITvDaiLy_DE_XUAT_SUA.docx',
        f1_edits, f1_party_start, f1_party, f1_notes)
process(UP+'ac061916-HopDong_DichVu_zVAS_KhachHang_DN_1__1.docx',
        '/home/user/finalproject/output/HopDong_DichVu_zVAS_KhachHangDN_DE_XUAT_SUA.docx',
        f2_edits, f2_party_start, f2_party, f2_notes)
