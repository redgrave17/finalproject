# -*- coding: utf-8 -*-
"""Slide mô hình luồng Dữ liệu Cá nhân Đại lý <-> GAPIT (zVAS/eSIM); eKYC = dữ liệu nhạy cảm."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
C=lambda h:RGBColor.from_string(h)
ZALO=C('0068FF');NAVY=C('001A40');BLUE2=C('003F88');ORANGE=C('F5862C');TEAL=C('0F6E56')
GREEN=C('22C55E');PURPLE=C('534AB7');RED=C('E2574C');GRAY=C('718096');GRAY2=C('4A5568')
LIGHT=C('F7F8F8');WHITE=C('FFFFFF');LN=C('E2E8F0');SLATE=C('94A3B8')
HEAD='Arial';BODY='Calibri';LOGO='/home/user/finalproject/assets/image4.png'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])
def rect(x,y,w,h,fill,line=None,rounded=True,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.shadow.inherit=False
    if fill is None: sh.fill.background()
    else: sh.fill.solid();sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line;sh.line.width=Pt(lw)
    return sh
def txt(x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
    tf.vertical_anchor=anchor;tf.margin_left=Inches(0.05);tf.margin_right=Inches(0.05);tf.margin_top=Inches(0.02);tf.margin_bottom=Inches(0.01)
    for i,(runs,sz,sb) in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(0);p.space_before=Pt(sb);p.line_spacing=1.02
        for t,c,b,fn in runs:
            r=p.add_run();r.text=t;r.font.size=Pt(sz);r.font.bold=b;r.font.color.rgb=c;r.font.name=fn
    return tb
def bullets(x,y,w,h,items,clr,sz=8.4,sa=3):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
    tf.margin_left=Inches(0.05);tf.margin_right=Inches(0.05);tf.margin_top=Inches(0.02)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(sa);p.line_spacing=1.0
        r=p.add_run();r.text='▪ ';r.font.size=Pt(sz);r.font.color.rgb=clr;r.font.bold=True;r.font.name=BODY
        for t,c,b in it:
            rr=p.add_run();rr.text=t;rr.font.size=Pt(sz);rr.font.color.rgb=c;rr.font.bold=b;rr.font.name=BODY
# top bar
rect(0,0,13.333,0.62,LIGHT,rounded=False)
s.shapes.add_picture(LOGO,Inches(0.45),Inches(0.12),height=Inches(0.40))
txt(1.78,0.12,3,0.40,[([('×  Zalo VAS',ZALO,True,HEAD)],14,0)],anchor=MSO_ANCHOR.MIDDLE)
rect(10.95,0.10,2.10,0.40,ORANGE);txt(10.95,0.10,2.10,0.40,[([('DATA FLOW MODEL',WHITE,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
txt(0.45,0.68,12.6,0.46,[([('Mô hình luồng Dữ liệu Cá nhân: Đại lý ↔ GAPIT (dịch vụ zVAS / eSIM)',NAVY,True,HEAD)],22,0)])
txt(0.45,1.16,12.6,0.26,[([('Các trường hợp Đại lý chuyển DLCN của người dùng cho GAPIT — phân loại Cơ bản / ',GRAY,False,BODY),('Nhạy cảm',RED,True,BODY),('; ',GRAY,False,BODY),('eKYC = dữ liệu cá nhân NHẠY CẢM',RED,True,BODY),('.',GRAY,False,BODY)],11,0)])
# actor chain
actors=[('Người dùng cuối / KH','Chủ thể dữ liệu',SLATE),('ĐẠI LÝ','Bên Kiểm soát',TEAL),
        ('GAPIT  (CMS)','Bên Xử lý',ZALO),('Zalo / VNG','Nhà phát hành',BLUE2),('Cơ quan nhà nước','Khi có yêu cầu',NAVY)]
albl=['thu thập','chuyển DLCN','eKYC / khai báo','cung cấp theo YC']
bw=2.20;step=2.49; y=1.58;h=0.64
for i,(t1,t2,clr) in enumerate(actors):
    x=0.45+i*step
    rect(x,y,bw,h,clr);txt(x,y+0.07,bw,0.30,[([(t1,WHITE,True,BODY)],10.5,0)],align=PP_ALIGN.CENTER)
    txt(x,y+0.36,bw,0.22,[([(t2,WHITE,False,BODY)],8,0)],align=PP_ALIGN.CENTER)
    if i<4:
        ax=x+bw
        txt(ax-0.05,y-0.16,step-bw+0.4,0.18,[([(albl[i],GRAY2,True,BODY)],7.3,0)],align=PP_ALIGN.CENTER)
        txt(ax,y,step-bw,h,[([('→',clr,True,HEAD)],20,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
# scenario cards
txt(0.45,2.42,12,0.22,[([('4 LUỒNG DỮ LIỆU CỤ THỂ',GRAY,True,BODY)],10,0)])
cards=[
 (TEAL,'1','Vận hành & Thống kê','Cơ bản / khử định danh',[
   [('Đại lý dùng CMS GAPIT để gắn Mã code zVAS vào SĐT người dùng, quản lý & duy trì tài khoản.',GRAY2,False)],
   [('Đại lý TỰ quản lý → ',GRAY2,False),('KHÔNG chuyển DLCN.',TEAL,True)],
   [('Cần dùng hệ thống GAPIT → dùng ',GRAY2,False),('Mã ID đã KHỬ ĐỊNH DANH',TEAL,True),(' (không còn là DLCN).',GRAY2,False)],
   [('Dữ liệu: SĐT / ID khử định danh.',GRAY2,False)]]),
 (ZALO,'2','CSKH / Khiếu nại / Hỗ trợ KT','Cơ bản',[
   [('Đại lý chuyển thông tin + nội dung sự việc cho GAPIT để xử lý ở cấp hệ thống.',GRAY2,False)],
   [('Dữ liệu: SĐT, email, mã đơn/mã code, nội dung khiếu nại.',GRAY2,False)],
   [('Vai trò: GAPIT = ',GRAY2,False),('Bên Xử lý',ZALO,True),(' theo chỉ dẫn của Đại lý.',GRAY2,False)]]),
 (RED,'3','Xác thực / eKYC   ⚠','DỮ LIỆU NHẠY CẢM',[
   [('Khách hàng không thể tự eKYC → Đại lý thu thập & chuyển ',GRAY2,False),('dữ liệu định danh',RED,True),(' cho GAPIT → GAPIT xử lý với Zalo.',GRAY2,False)],
   [('Dữ liệu: ảnh CCCD/giấy tờ định danh (và sinh trắc nếu có) — ',GRAY2,False),('NHẠY CẢM',RED,True),('.',GRAY2,False)],
   [('Yêu cầu: ',RED,True),('đồng ý riêng + mã hóa/bảo mật nghiêm ngặt + tối thiểu hóa + thông báo xử lý DLCN nhạy cảm.',GRAY2,False)]]),
 (NAVY,'4','Cung cấp cho cơ quan nhà nước','Nghĩa vụ pháp lý',[
   [('Khi có hành vi gian lận → GAPIT gửi lại dữ liệu cho Zalo/VNG để khai báo & cung cấp cho cơ quan có thẩm quyền.',GRAY2,False)],
   [('Cơ sở: ',GRAY2,False),('nghĩa vụ pháp lý / yêu cầu của cơ quan có thẩm quyền',NAVY,True),(' (không cần đồng ý bổ sung).',GRAY2,False)]]),
]
cw=3.05;cx=[0.45,3.62,6.79,9.96];cy=2.68;chh=3.30
for (clr,num,title,tag,items),x in zip(cards,cx):
    rect(x,cy,cw,chh,WHITE,line=clr,lw=1.5)
    rect(x,cy,cw,0.74,clr)
    txt(x+0.14,cy+0.06,cw-0.28,0.30,[([(num+' · '+title,WHITE,True,BODY)],10,0)])
    txt(x+0.14,cy+0.42,cw-0.28,0.26,[([(tag,WHITE,True,BODY)],8.6,0)])
    bullets(x+0.08,cy+0.84,cw-0.16,chh-0.9,items,clr,sz=8.5,sa=4)
# note bar
rect(0.45,6.12,12.45,0.82,C('FDEEEC'),line=RED,lw=1.25)
txt(0.62,6.16,12.15,0.74,[
 ([('Phân loại & vai trò: ',NAVY,True,BODY),('Đại lý = Bên Kiểm soát (thu thập, quan hệ trực tiếp với người dùng) · GAPIT = Bên Xử lý (xử lý cấp hệ thống theo chỉ dẫn).',GRAY2,False,BODY)],9,0),
 ([('⚠ Riêng eKYC là DỮ LIỆU NHẠY CẢM',RED,True,BODY),(' → áp dụng sự đồng ý RIÊNG + biện pháp bảo mật nghiêm ngặt + thông báo xử lý dữ liệu nhạy cảm theo Luật BVDLCN 91/2025 & NĐ 356/2025. Nguyên tắc khử định danh/tối thiểu hóa cho luồng vận hành.',GRAY2,False,BODY)],9,2),
])
# footer
rect(0.45,7.20,12.45,0.22,BLUE2);txt(0.6,7.20,12.1,0.22,[([('GAPIT JSC  |  gapit.com.vn',WHITE,True,BODY),('     —  Mô hình luồng DLCN Đại lý ↔ GAPIT (phục vụ DPA & rà soát Pháp chế)',WHITE,False,BODY)],8,0)],anchor=MSO_ANCHOR.MIDDLE)
import os;os.makedirs('/home/user/finalproject/output',exist_ok=True)
out='/home/user/finalproject/output/GAPIT_MoHinh_LuongDLCN_DaiLy.pptx';prs.save(out);print('SAVED',out)
