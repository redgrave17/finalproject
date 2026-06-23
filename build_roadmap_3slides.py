# -*- coding: utf-8 -*-
"""Tách Business Roadmap zVAS thành 3 slide riêng:
 1) Roadmap Now/Next/Later  2) Biên lợi nhuận GAPIT & ĐL cấp II  3) Order Flow zVAS CMS (swimlane).
Branding GAPIT/Zalo."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

C=lambda h:RGBColor.from_string(h)
ZALO=C('0068FF'); NAVY=C('001A40'); BLUE2=C('003F88'); ORANGE=C('F5862C')
TEAL=C('0F6E56'); GREEN=C('22C55E'); PURPLE=C('534AB7'); AMBER=C('BA7517')
CLBLUE=C('185FA5'); GRAY=C('718096'); GRAY2=C('4A5568'); LIGHT=C('F7F8F8')
WHITE=C('FFFFFF'); LN=C('E2E8F0'); SLATE=C('94A3B8')
HEAD='Arial'; BODY='Calibri'
LOGO='/home/user/finalproject/assets/image4.png'

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def rect(sl,x,y,w,h,fill,line=None,rounded=False,lw=0.75):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Inches(x),Inches(y),Inches(w),Inches(h))
    sh.shadow.inherit=False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    return sh

def txt(sl,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.02); tf.margin_bottom=Inches(0.01)
    for i,(runs,sz,sb) in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(0); p.space_before=Pt(sb); p.line_spacing=1.03
        for t,c,b,fn in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c; r.font.name=fn
    return tb

def bullets(sl,x,y,w,h,items,clr,sz=11,sa=5):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.02)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(sa); p.space_before=Pt(0); p.line_spacing=1.02
        r=p.add_run(); r.text='▪ '; r.font.size=Pt(sz); r.font.color.rgb=clr; r.font.bold=True; r.font.name=BODY
        r2=p.add_run(); r2.text=it; r2.font.size=Pt(sz); r2.font.color.rgb=GRAY2; r2.font.name=BODY

def frame(sl,chip,title,subtitle_paras):
    rect(sl,0,0,13.333,0.62,LIGHT)
    sl.shapes.add_picture(LOGO,Inches(0.45),Inches(0.12),height=Inches(0.40))
    txt(sl,1.78,0.12,3.0,0.40,[([('×  Zalo VAS',ZALO,True,HEAD)],14,0)],anchor=MSO_ANCHOR.MIDDLE)
    rect(sl,10.95,0.10,2.10,0.40,ORANGE,rounded=True)
    txt(sl,10.95,0.10,2.10,0.40,[([(chip,WHITE,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(sl,0.45,0.70,12.6,0.46,[([(title,NAVY,True,HEAD)],25,0)])
    txt(sl,0.45,1.22,12.6,0.30,subtitle_paras)
    rect(sl,0.45,7.20,12.45,0.22,BLUE2,rounded=True)
    txt(sl,0.6,7.20,12.1,0.22,[([('GAPIT JSC  |  gapit.com.vn',WHITE,True,BODY),
       ('     —  Business Roadmap sản phẩm zVAS',WHITE,False,BODY)],8,0)],anchor=MSO_ANCHOR.MIDDLE)

blank=prs.slide_layouts[6]

# ============================================================ SLIDE 1: ROADMAP
s1=prs.slides.add_slide(blank)
frame(s1,'ROADMAP · 1/3','Business Roadmap zVAS — Now · Next · Later',
 [([('Roadmap theo kết quả/outcome & horizon — không cố định ngày, ưu tiên theo chủ đề (theme) và giá trị.',GRAY,False,BODY)],11,0),
  ([('Khung tham khảo: Adrenalin — “Roadmapping Your Digital Product: A Definitive Guide”.',GRAY,False,BODY)],9,1)])
cols=[
 (ZALO,'NOW','Q2–Q3 / 2026','Go-live & Validate','Outcome: vận hành B2B ổn định · dòng tiền trả trước · khởi động mạng đại lý',
  ['Hoàn thiện zVAS CMS: catalog, báo giá (prorate + VAT), hóa đơn, kích hoạt',
   'Code pool & cảnh báo hạn kích hoạt 6/3/2/1 tháng (trong 12 tháng từ nhập kho)',
   'Go-live KHDN + ký HĐ Đại lý cấp II; DPA theo mẫu PDPA(VNI)',
   'Chuẩn hóa cơ chế trả trước + bảng giá & chiết khấu theo bậc',
   'Đào tạo CSKH cấp 1; quy trình đổi Mã code lỗi (≤ 24h)']),
 (TEAL,'NEXT','Q4/2026 – Q1/2027','Scale & Enable','Outcome: nhân rộng kênh đại lý · tăng MRR · giảm thao tác thủ công',
  ['Mở rộng Đại lý cấp II: RACI, onboarding, KPI doanh số tối thiểu',
   'Portal Đại lý / Affiliate; báo cáo MRR & tồn kho Mã code',
   'Tự động đối soát CSV zbox.vn; auto-select code sắp hết hạn',
   'Thêm bundle (zBusiness + zCloud + zStyle)',
   'Khuyến mại có kiểm soát giá sàn công khai (VNG)']),
 (PURPLE,'LATER','2027 +','Optimize & Expand','Outcome: tối ưu biên lợi nhuận · mở rộng hệ sinh thái · tích hợp sâu',
  ['API hóa khi VNG mở API → loại bỏ thao tác CSV thủ công',
   'Cross-sell eSIM / SIM doanh nghiệp; combo đa dịch vụ',
   'Analytics hành vi & dự báo nhu cầu',
   'Tối ưu pricing & biên theo bậc số lượng',
   'Loyalty cho Đại lý & KHDN; mở rộng vùng/ngành']),
]
xs=[0.45,4.64,8.83]; cw=4.05
for (clr,hz,tf_,tag,oc,items),x in zip(cols,xs):
    rect(s1,x,1.70,cw,0.92,clr,rounded=True)
    txt(s1,x+0.18,1.76,cw-0.34,0.34,[([(hz,WHITE,True,HEAD),('    '+tf_,WHITE,False,BODY)],16,0)])
    txt(s1,x+0.18,2.18,cw-0.34,0.34,[([(tag,WHITE,True,BODY)],12,0)])
    rect(s1,x,2.70,cw,3.95,WHITE,line=LN)
    rect(s1,x,2.70,cw,0.62,C('F4F7FB'))
    txt(s1,x+0.14,2.74,cw-0.28,0.56,[([(oc,clr,True,BODY)],9.5,0)])
    bullets(s1,x+0.10,3.42,cw-0.2,3.15,items,clr,sz=11,sa=7)

# ============================================================ SLIDE 2: MARGINS
s2=prs.slides.add_slide(blank)
frame(s2,'BIÊN LỢI NHUẬN · 2/3','Biên Lợi Nhuận — GAPIT & Đại Lý Cấp II',
 [([('Tổng biên kênh ≈ 20% (chênh giá bán lẻ Zalo vs giá VNG cấp cho GAPIT) · GAPIT giữ = 20% − chiết khấu cấp cho Đại lý cấp II.',GRAY,False,BODY)],11,0)])
# Table
tx,ty=0.45,1.78; colw=[2.05,1.85,1.7,1.95,1.35]; rh=0.5
hdr=['Bậc SL (code/tháng)','CK Đại lý cấp II','Biên Đại lý II','Biên GAPIT giữ','Tổng biên kênh']
rows=[['50 – 99','5%','5%','15%','20%'],['100 – 299','8%','8%','12%','20%'],['≥ 300','10%','10%','10%','20%']]
cx=tx
for j,h in enumerate(hdr):
    rect(s2,cx,ty,colw[j],rh,BLUE2); txt(s2,cx,ty,colw[j],rh,[([(h,WHITE,True,BODY)],10,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE); cx+=colw[j]
for i,row in enumerate(rows):
    cx=tx; yy=ty+rh*(i+1)
    for j,val in enumerate(row):
        fill=WHITE if i%2==0 else C('F4F7FB')
        if j==3: fill=C('FFF1E6')
        rect(s2,cx,yy,colw[j],rh,fill,line=LN,lw=0.5)
        col=ORANGE if j==3 else (NAVY if j==0 else GRAY2); bold=(j in (0,3))
        txt(s2,cx,yy,colw[j],rh,[([(val,col,bold,BODY)],11,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE); cx+=colw[j]
tw=sum(colw)
# Stacked bars
txt(s2,0.45,3.66,tw,0.26,[([('Cấu trúc biên trên 1 đơn vị (× 100% giá bán lẻ Zalo)',GRAY,True,BODY)],10,0)])
bars=[('Bậc 50–99',80,15,5),('Bậc 100–299',80,12,8),('Bậc ≥300',80,10,10)]
bx=0.45; barw=tw-1.4; by=4.0
for i,(lbl,vng,gap,ag) in enumerate(bars):
    yy=by+i*0.52
    txt(s2,bx,yy,1.3,0.42,[([(lbl,NAVY,True,BODY)],9.5,0)],anchor=MSO_ANCHOR.MIDDLE)
    x0=bx+1.4
    segs=[('Giá VNG '+str(vng)+'%',vng,SLATE,WHITE),('GAPIT '+str(gap)+'%',gap,ORANGE,WHITE),('ĐL '+str(ag)+'%',ag,GREEN,WHITE)]
    cxx=x0
    for name,pct,fc,tc in segs:
        wseg=barw*pct/100.0
        rect(s2,cxx,yy,wseg,0.42,fc,rounded=False)
        txt(s2,cxx,yy,wseg,0.42,[([(name,tc,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        cxx+=wseg
# Right: examples
exx=10.0
rect(s2,exx,1.78,2.9,0.4,NAVY,rounded=True)
txt(s2,exx,1.78,2.9,0.4,[([('VÍ DỤ — zBusiness 12 tháng',WHITE,True,BODY)],10,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
txt(s2,exx,2.26,2.9,0.5,[([('Zalo retail 1.990.000đ',NAVY,True,BODY)],10.5,0),
   ([('VNG → GAPIT 1.592.000đ (−20%)',GRAY2,False,BODY)],9.5,1)])
def excard(y,clr,tier,buy,gp,gpp,ag,agp):
    rect(s2,exx,y,2.9,1.18,WHITE,line=clr,rounded=True,lw=1.25); rect(s2,exx,y,0.10,1.18,clr)
    txt(s2,exx+0.2,y+0.06,2.6,0.26,[([(tier,clr,True,BODY)],10,0)])
    txt(s2,exx+0.2,y+0.34,2.6,0.8,[
       ([('ĐL mua: ',GRAY2,False,BODY),(buy,NAVY,True,BODY)],9.3,0),
       ([('Biên GAPIT: ',GRAY2,False,BODY),(gp+' ('+gpp+')',ORANGE,True,BODY)],9.3,2),
       ([('Biên Đại lý: ',GRAY2,False,BODY),(ag+' ('+agp+')',GREEN,True,BODY)],9.3,2)])
excard(2.86,GREEN,'Bậc ≥300 · CK 10%','1.791.000đ','199.000đ','10%','199.000đ','10%')
excard(4.18,ORANGE,'Bậc 50–99 · CK 5%','1.890.500đ','298.500đ','15%','99.500đ','5%')
# note
rect(s2,0.45,5.95,12.45,0.95,C('FFF6E9'),line=ORANGE,rounded=True)
txt(s2,0.65,6.02,12.1,0.85,[
 ([('Lưu ý: ',ORANGE,True,BODY),('“Tổng biên kênh ≈20%” suy ra từ giá zBusiness 12 tháng tại Phụ lục HĐ VNG–GAPIT; các gói khác (zCloud/zStyle…) theo Phụ lục tương ứng — cập nhật % thực tế khi có giá từng gói.',GRAY2,False,BODY)],9.5,0),
 ([('Chiết khấu Đại lý cấp II tính trên Giá bán lẻ trực tiếp của Zalo (giá sàn công khai). Mã code đã mua: không hoàn/đổi (chính sách VNG). Đại lý không được bán công khai dưới giá sàn.',GRAY2,False,BODY)],9.5,3)])

# ============================================================ SLIDE 3: ORDER FLOW (swimlane)
s3=prs.slides.add_slide(blank)
frame(s3,'ORDER FLOW · 3/3','Order Flow trên zVAS CMS (6 bước · không GapOne)',
 [([('Phân luồng theo các bên tham gia trên hệ thống zVAS CMS (Service Platform) — quản lý code pool, đơn hàng, kích hoạt, hóa đơn.',GRAY,False,BODY)],11,0)])
LX=0.45; LW=1.45; gap=0.06
n=6; CW=(13.05-LX-LW-(n)*gap)/n
def colx(i): return LX+LW+gap+i*(CW+gap)
hy=1.66; hh=0.5
steps=['View Catalog','Chọn Bundle / Seats','Báo giá (Prorate + VAT)','Thanh toán + Hóa đơn','Subscription Active','Pool & Phân bổ seats']
# step headers
for i,st in enumerate(steps):
    rect(s3,colx(i),hy,CW,hh,BLUE2,rounded=True)
    txt(s3,colx(i),hy+0.02,CW,0.20,[([('0'+str(i+1),C('93C5FD'),True,HEAD)],11,0)],align=PP_ALIGN.CENTER)
    txt(s3,colx(i),hy+0.22,CW,0.26,[([(st,WHITE,True,BODY)],7.8,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
# actor rows
rowy=2.22; rh=0.98
actors=[('Client','Owner / Admin',CLBLUE,C('E6F1FB')),
        ('zVAS CMS','System · auto',PURPLE,C('EEEDFE')),
        ('GAPIT Sales/Admin','Internal',TEAL,C('E1F5EE')),
        ('Payment / Ngoài','ZaloPay·Bank·VNPT',AMBER,C('FAEEDA'))]
cells=[
 # Client
 [['Duyệt catalog zBusiness/zCloud/zStyle/SIM','So sánh bundle & bảng giá'],
  ['Chọn bundle + số seat','Chu kỳ tháng/năm · xem giá tạm tính'],
  ['Review quote PDF','Xác nhận prorate · approve đơn'],
  ['Chuyển khoản / ZaloPay','Upload payment proof'],
  ['Nhận email active','Truy cập Client Portal'],
  ['Vào Pool · phân bổ seat','Theo phòng ban / org tree']],
 # CMS
 [['Render catalog theo phân quyền','Pricing rules & feature matrix'],
  ['Validate seat (min/max)','Cảnh báo code hạn 6/3/2/1 thg · tạo order (draft)'],
  ['Tính prorate · VAT 10%','Generate quote PDF'],
  ['Xuất invoice PDF (VAT·MST)','Set pending_payment'],
  ['Tạo subscription · set dates','Auto-select code sắp hết hạn · draft→active'],
  ['Cộng seat vào Pool','available=purchased−allocated · log']],
 # Sales/Admin
 [['Demo catalog cho KH','Khởi tạo order hộ (sales-assist)'],
  ['Cấu hình custom bundle','Enterprise pricing · note nội bộ'],
  ['Approve quote custom','Gửi KH qua email / Zalo'],
  ['⚠ Verify thanh toán thủ công','Mark paid → CMS kích hoạt'],
  ['Cập nhật KPI pipeline','Ghi nhận MRR mới'],
  ['Theo dõi tồn kho code pool','Chuẩn bị activation']],
 # External
 [['—'],['—'],['—'],
  ['ZaloPay/Bank xử lý · transaction ID','e-Contract bởi VNPT'],
  ['—'],['—']],
]
for r,(an,asub,clr,bg) in enumerate(actors):
    yy=rowy+r*rh
    rect(s3,LX,yy,LW,rh,clr,rounded=True)
    txt(s3,LX+0.06,yy+0.10,LW-0.12,0.5,[([(an,WHITE,True,BODY)],9.5,0),([(asub,WHITE,False,BODY)],7.5,1)],anchor=MSO_ANCHOR.MIDDLE)
    for i in range(6):
        items=cells[r][i]
        empty=(items==['—'])
        rect(s3,colx(i),yy,CW,rh,C('F1F5F9') if empty else bg,line=LN,lw=0.5,rounded=True)
        if empty:
            txt(s3,colx(i),yy,CW,rh,[([('—',SLATE,True,BODY)],14,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        else:
            bullets(s3,colx(i)+0.02,yy+0.04,CW-0.05,rh-0.06,items,clr,sz=7.3,sa=2)
# status row
sy=rowy+4*rh+0.06
txt(s3,LX,sy,LW,0.3,[([('Trạng thái',GRAY,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
stat=[('—',C('F1F5F9'),GRAY),('draft',C('EEF2FF'),C('4338CA')),('draft (quoted)',C('EEF2FF'),C('4338CA')),
      ('pending_payment',C('FEF3C7'),C('92400E')),('active ✓',C('DCFCE7'),C('14532D')),('active + pool',C('DCFCE7'),C('14532D'))]
for i,(t,bg,fc) in enumerate(stat):
    rect(s3,colx(i),sy,CW,0.3,bg,rounded=True)
    txt(s3,colx(i),sy,CW,0.3,[([(t,fc,True,BODY)],8,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

import os
os.makedirs('/home/user/finalproject/output',exist_ok=True)
out='/home/user/finalproject/output/GAPIT_zVAS_Business_Roadmap_3slides.pptx'
prs.save(out); print('SAVED',out)
