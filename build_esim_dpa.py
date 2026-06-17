# -*- coding: utf-8 -*-
"""Build the eSIM Data Processing Agreement (Vietnamese), modeled on the WinX DPA core,
referencing the GAPIT eSIM Agent Contract (HĐ Đại lý số …/2026/HĐĐL/GPT)."""
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = '/home/user/finalproject/output/Thoa_thuan_Xu_ly_DLCN_eSIM_GAPIT_DaiLy.docx'
d = docx.Document()

# Base style
st = d.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)

def H(text, size=13, center=False, space_before=10):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    return p

def P(text, just=True, italic=False, bold=False, indent=None):
    p = d.add_paragraph()
    r = p.add_run(text); r.italic = italic; r.bold = bold
    if just: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(4)
    return p

def NOTE(text):
    p = d.add_paragraph()
    r = p.add_run(text); r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.italic=True; r.font.size=Pt(11)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

# ---------- Title ----------
H('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', size=12, center=True, space_before=0)
H('Độc lập – Tự do – Hạnh phúc', size=12, center=True, space_before=0)
P('--------o0o--------', just=False).alignment = WD_ALIGN_PARAGRAPH.CENTER
H('THỎA THUẬN XỬ LÝ DỮ LIỆU CÁ NHÂN', size=15, center=True)
H('(Áp dụng cho dịch vụ SIM/eSIM du lịch phân phối qua Đại lý)', size=11, center=True, space_before=0)
P('Số: ……/2026/DPA/GPT', just=False).alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- Caveat note (ties to reviewer comment 3) ----------
NOTE('LƯU Ý PHÁP LÝ (cần Bộ phận Pháp chế xác nhận trước khi ký): Văn bản này được dự thảo trên cơ sở "lõi" của '
     'Thỏa thuận Xử lý Dữ liệu Cá nhân (DPA) GAPIT–WinX và gắn với Hợp đồng Đại lý phân phối SIM/eSIM du lịch. '
     'Việc xác định vai trò Bên Kiểm soát/Bên Xử lý phụ thuộc vào luồng dữ liệu thực tế của mô hình B2B (theo góp ý của maiht tại Điều 8 Hợp đồng Đại lý). '
     'Dự thảo này mặc định: Đại lý (Bên B) là BÊN KIỂM SOÁT DỮ LIỆU (thu thập, có quan hệ trực tiếp với khách hàng cuối, bán nhân danh chính mình theo Điều 1.2 Hợp đồng Đại lý) và GAPIT (Bên A) là BÊN XỬ LÝ DỮ LIỆU (xử lý đơn hàng trên hệ thống CMS để cấp eSIM, gửi mã QR, xuất hóa đơn, hỗ trợ kỹ thuật). '
     'Riêng đối với (i) việc xuất hóa đơn GTGT nhân danh GAPIT và (ii) việc chuyển dữ liệu cho nhà cung cấp eSIM ở nước ngoài (Consortio), GAPIT có thể đóng vai trò Bên Kiểm soát độc lập — trường hợp này cân nhắc dùng "thỏa thuận chia sẻ dữ liệu giữa các bên kiểm soát độc lập" thay cho mô hình kiểm soát–xử lý.')

# ---------- Preamble ----------
P('Thỏa thuận Xử lý Dữ liệu Cá nhân này (“Thỏa thuận”) được lập và có hiệu lực kể từ ngày …… tháng …… năm 20…… '
  '(“Ngày Hiệu Lực”) giữa các bên sau:')
P('1. BÊN KIỂM SOÁT DỮ LIỆU (BÊN B – ĐẠI LÝ): ………………………………………………………………, '
  'một tổ chức được thành lập và hoạt động hợp pháp theo pháp luật Việt Nam, có trụ sở chính tại '
  '……………………………………………………… (sau đây gọi là “Bên Kiểm Soát Dữ Liệu”); và')
P('2. BÊN XỬ LÝ DỮ LIỆU (BÊN A): CÔNG TY CỔ PHẦN GAPIT, một công ty được thành lập và hoạt động hợp pháp theo '
  'pháp luật Việt Nam, có trụ sở chính tại Phòng 902, Tầng 9, D10 Giảng Võ, Phường Giảng Võ, Thành phố Hà Nội, '
  'Việt Nam; Mã số thuế: 0101816595 (sau đây gọi là “Bên Xử Lý Dữ Liệu”).')
P('Bên Kiểm Soát Dữ Liệu và Bên Xử Lý Dữ Liệu sau đây được gọi riêng là “Bên” và gọi chung là “Các Bên”.')

H('CÁC CĂN CỨ', size=12)
P('(A) Các Bên đã ký kết Hợp đồng Đại lý phân phối Sản phẩm SIM/eSIM du lịch số ……/2026/HĐĐL/GPT '
  '(“Hợp Đồng Cơ Sở”), theo đó Bên Kiểm Soát Dữ Liệu làm đại lý bán Sản phẩm SIM/eSIM du lịch do Bên Xử Lý Dữ Liệu '
  'phân phối cho khách hàng cuối, đặt hàng và thanh toán qua hệ thống CMS của Bên Xử Lý Dữ Liệu, trong đó phát sinh '
  'các hoạt động xử lý dữ liệu cá nhân của khách hàng cuối.')
P('(B) Liên quan đến việc xử lý dữ liệu cá nhân theo Hợp Đồng Cơ Sở và Thỏa thuận này, Các Bên thống nhất rằng '
  'Bên Kiểm Soát Dữ Liệu là bên kiểm soát dữ liệu và Bên Xử Lý Dữ Liệu là bên xử lý dữ liệu theo quy định của '
  'pháp luật về bảo vệ dữ liệu cá nhân hiện hành.')
P('(C) Các Bên mong muốn thiết lập một thỏa thuận xử lý dữ liệu cá nhân nhằm bảo đảm việc xử lý dữ liệu cá nhân '
  'tuân thủ các quy định của pháp luật về bảo vệ dữ liệu cá nhân.')
P('VÌ VẬY, các Bên thống nhất như sau:', bold=True)

# ---------- Article 1 ----------
H('Điều 1. Định nghĩa và giải thích')
P('Trong Thỏa thuận này, các thuật ngữ dưới đây được hiểu như sau:')
defs = [
 ('“Cơ quan có thẩm quyền”', 'là bất kỳ cơ quan nhà nước có thẩm quyền nào của Việt Nam, bao gồm Cục An ninh mạng và '
  'Phòng, chống tội phạm sử dụng công nghệ cao thuộc Bộ Công an Việt Nam và các cơ quan tư pháp có thẩm quyền giám sát, '
  'thanh tra, xét xử việc tuân thủ quy định về bảo vệ dữ liệu cá nhân và áp dụng các quy định pháp luật có liên quan.'),
 ('“Vi phạm dữ liệu”', 'là việc không tuân thủ các Quy định về bảo vệ dữ liệu cá nhân, bao gồm việc mất mát, rò rỉ, '
  'sử dụng trái phép hoặc bất hợp pháp, sao chép, sửa đổi, tiết lộ, phá hủy hoặc truy cập dữ liệu cá nhân, hoặc bất kỳ '
  'hành vi nào khác gây ảnh hưởng tiêu cực đến dữ liệu cá nhân được chuyển giao theo Thỏa thuận này.'),
 ('“Quy định về bảo vệ dữ liệu cá nhân”', 'là bất kỳ luật, đạo luật, quyết định, nghị định, chỉ thị, văn bản lập pháp, '
  'lệnh, quy định, quy tắc hoặc hạn chế có tính ràng buộc nào của Việt Nam, bao gồm nhưng không giới hạn ở Luật Bảo vệ '
  'Dữ liệu Cá nhân số 91/2025/QH15 được Quốc hội ban hành ngày 26 tháng 6 năm 2025 và Nghị định số 356/2025/NĐ-CP '
  'quy định chi tiết một số điều và biện pháp thi hành Luật Bảo vệ Dữ liệu Cá nhân do Chính phủ ban hành ngày 31 tháng 12 '
  'năm 2025, được sửa đổi, bổ sung theo từng thời kỳ, áp dụng đối với việc xử lý dữ liệu cá nhân.'),
 ('“Chủ thể dữ liệu cá nhân”', 'là cá nhân mà dữ liệu cá nhân phản ánh hoặc liên quan đến, trong phạm vi Thỏa thuận này '
  'chủ yếu là khách hàng cuối mua Sản phẩm SIM/eSIM du lịch thông qua Bên Kiểm Soát Dữ Liệu.'),
 ('“Dữ liệu cá nhân”', 'là dữ liệu hoặc thông tin ở dạng số hoặc các hình thức khác có thể xác định hoặc hỗ trợ việc '
  'xác định một cá nhân cụ thể, bao gồm dữ liệu cá nhân cơ bản và dữ liệu cá nhân nhạy cảm. Dữ liệu cá nhân đã được ẩn danh '
  'sẽ không còn được coi là dữ liệu cá nhân.'),
 ('“Nhân sự”', 'là bất kỳ người lao động nào, bao gồm người quản lý và nhân viên của Bên Xử Lý Dữ Liệu (bao gồm cả người '
  'làm việc theo hợp đồng lao động bán thời gian, không xác định thời hạn hoặc xác định thời hạn) có quyền truy cập và được '
  'Bên Xử Lý Dữ Liệu phân công thực hiện việc xử lý dữ liệu cá nhân.'),
 ('“Xử lý dữ liệu cá nhân”', 'là bất kỳ hoạt động nào tác động đến dữ liệu cá nhân, bao gồm một hoặc nhiều hoạt động như: '
  'thu thập, phân tích, tổng hợp, mã hóa, giải mã, chỉnh sửa, xóa, hủy, ẩn danh, cung cấp, tiết lộ, chuyển giao dữ liệu cá nhân '
  'hoặc các hoạt động khác có liên quan đến dữ liệu cá nhân.'),
 ('“Mục đích”', 'là mục đích xử lý dữ liệu cá nhân được quy định tại Điều 2 của Thỏa thuận này.'),
 ('“Hệ thống CMS”', 'là hệ thống quản lý đặt hàng, thanh toán, xuất hóa đơn và tra cứu do Bên Xử Lý Dữ Liệu xây dựng và '
  'vận hành để phục vụ việc phân phối Sản phẩm SIM/eSIM du lịch theo Hợp Đồng Cơ Sở.'),
]
for term, body in defs:
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(term + ' '); r.bold = True
    p.add_run(body)
P('Trong Thỏa thuận này, (a) các từ “bao gồm” được hiểu là “bao gồm nhưng không giới hạn”; (b) từ ở dạng số ít bao gồm cả '
  'dạng số nhiều và ngược lại; và (c) tiêu đề các điều khoản không ảnh hưởng đến việc giải thích Thỏa thuận này.')

# ---------- Article 2 ----------
H('Điều 2. Mục đích và Phương thức Xử lý Dữ liệu Cá nhân')
P('2.1. Trong quá trình thực hiện Hợp Đồng Cơ Sở, Bên Xử Lý Dữ Liệu, thay mặt Bên Kiểm Soát Dữ Liệu, sẽ thực hiện các hoạt '
  'động Xử Lý đối với Dữ Liệu Cá Nhân của khách hàng cuối nhằm: (i) tiếp nhận và xử lý đơn hàng trên Hệ thống CMS; '
  '(ii) cấp Sản phẩm và gửi mã QR/thông tin kích hoạt eSIM cho khách hàng cuối qua email; (iii) xuất hóa đơn giá trị gia tăng '
  'tương ứng với đơn hàng đã thanh toán; (iv) hỗ trợ kỹ thuật, hướng dẫn kích hoạt và xử lý khiếu nại; và (v) lưu trữ, tra cứu, '
  'đối soát đơn hàng theo Hợp Đồng Cơ Sở (“Mục đích”).')
P('2.2. Việc Xử Lý Dữ Liệu Cá Nhân cho bất kỳ mục đích nào khác ngoài Mục đích chỉ được thực hiện khi có chỉ dẫn bằng văn bản '
  'của Bên Kiểm Soát Dữ Liệu hoặc theo yêu cầu của Quy định về Bảo vệ Dữ liệu cá nhân.')
P('2.3. Bên Xử Lý Dữ Liệu không được Xử Lý Dữ Liệu Cá Nhân bằng các phương thức khác với các phương thức đã được Bên Kiểm Soát '
  'Dữ Liệu thông báo và chấp thuận trước đó, trừ khi có chỉ dẫn khác bằng văn bản của Bên Kiểm Soát Dữ Liệu hoặc được cho phép '
  'theo Quy định về Bảo vệ Dữ liệu cá nhân.')

# ---------- Article 3 ----------
H('Điều 3. Nghĩa vụ Cơ bản của Bên Xử Lý Dữ Liệu')
items3 = [
 '3.1. Bên Xử Lý Dữ Liệu phải triển khai đầy đủ và phù hợp các biện pháp kỹ thuật, hành chính và tổ chức, đồng thời tiến hành '
 'kiểm tra định kỳ để đánh giá hiệu quả của các biện pháp đó nhằm bảo vệ Dữ Liệu Cá Nhân. Các biện pháp này phải tuân thủ các '
 'tiêu chuẩn phù hợp với các yêu cầu áp dụng đối với Bên Kiểm Soát Dữ Liệu cũng như các Quy định về Bảo vệ Dữ liệu cá nhân.',
 '3.2. Bên Xử Lý Dữ Liệu phải thông báo cho Bên Kiểm Soát Dữ Liệu, trong vòng 24 (hai mươi bốn) giờ kể từ khi phát hiện, về bất '
 'kỳ vi phạm nào đối với Quy định về Bảo vệ Dữ liệu cá nhân (bao gồm bất kỳ Vi phạm Dữ liệu nào) liên quan đến Dữ Liệu Cá Nhân, '
 'và theo yêu cầu của Bên Kiểm Soát Dữ Liệu, cung cấp đầy đủ thông tin và bằng chứng để Bên Kiểm Soát Dữ Liệu có thể thực hiện '
 'các nghĩa vụ liên quan, đồng thời phản hồi theo chỉ dẫn của Bên Kiểm Soát Dữ Liệu.',
 '3.3. Khi xảy ra trường hợp quy định tại Điều 3.2, theo yêu cầu bằng văn bản của Bên Kiểm Soát Dữ Liệu, Bên Xử Lý Dữ Liệu phải '
 'hỗ trợ Bên Kiểm Soát Dữ Liệu trong việc báo cáo hoặc tự mình báo cáo sự tồn tại của vi phạm cho Cơ quan có thẩm quyền, thực hiện '
 'các biện pháp theo Quy định về Bảo vệ Dữ liệu cá nhân và phối hợp với Bên Kiểm Soát Dữ Liệu cũng như Cơ quan có thẩm quyền để '
 'điều tra, giảm thiểu và khắc phục từng Vi phạm Dữ liệu.',
 '3.4. Bên Xử Lý Dữ Liệu phải nhanh chóng thông báo cho Bên Kiểm Soát Dữ Liệu về bất kỳ yêu cầu hoặc thắc mắc nào nhận được từ '
 'Chủ thể dữ liệu cá nhân hoặc từ Cơ quan có thẩm quyền liên quan đến Dữ Liệu Cá Nhân và hỗ trợ Bên Kiểm Soát Dữ Liệu trong việc '
 'phản hồi (ví dụ: ngừng Xử Lý hoặc xóa Dữ Liệu Cá Nhân theo yêu cầu).',
 '3.5. Bên Xử Lý Dữ Liệu chỉ được Xử Lý Dữ Liệu Cá Nhân một cách nghiêm ngặt theo đúng Thỏa thuận này và chịu trách nhiệm đối với '
 'Bên Kiểm Soát Dữ Liệu về bất kỳ thiệt hại nào phát sinh từ việc xử lý dữ liệu cá nhân của mình.',
 '3.6. Bên Xử Lý Dữ Liệu phải cung cấp mọi hỗ trợ cần thiết cho Bên Kiểm Soát Dữ Liệu trong việc thực hiện các nghĩa vụ của Bên '
 'Kiểm Soát Dữ Liệu theo Quy định về Bảo vệ Dữ liệu cá nhân, và thực hiện đầy đủ mọi nghĩa vụ áp dụng đối với mình với tư cách là '
 'bên xử lý dữ liệu.',
 '3.7. Bên Xử Lý Dữ Liệu không được yêu cầu Bên Kiểm Soát Dữ Liệu thanh toán bất kỳ khoản tiền nào liên quan đến việc Xử Lý Dữ '
 'Liệu Cá Nhân theo Thỏa thuận này.',
]
for it in items3: P(it)

# ---------- Article 4 ----------
H('Điều 4. Thuê Bên Xử Lý Phụ')
P('4.1. Bên Xử Lý Dữ Liệu không được thuê, giao khoán hoặc ủy thác toàn bộ hoặc một phần việc Xử Lý Dữ Liệu theo Thỏa thuận này '
  'cho một bên xử lý phụ, trừ khi tuân thủ các yêu cầu tại Điều 4.2 và bên xử lý phụ đã thiết lập, triển khai các biện pháp bảo vệ '
  'Dữ Liệu Cá Nhân ở mức tương đương hoặc cao hơn so với Bên Xử Lý Dữ Liệu, đáp ứng các tiêu chí do Bên Kiểm Soát Dữ Liệu thiết lập.')
P('4.2. Trước khi thực hiện việc xử lý phụ, Bên Xử Lý Dữ Liệu phải cung cấp cho Bên Kiểm Soát Dữ Liệu báo cáo bằng văn bản nêu rõ '
  'thông tin chi tiết về bên xử lý phụ (bao gồm phương thức xử lý và biện pháp bảo vệ Dữ Liệu Cá Nhân) và phải nhận được sự chấp '
  'thuận của Bên Kiểm Soát Dữ Liệu.')
P('4.3. Sau khi được chấp thuận, Bên Xử Lý Dữ Liệu phải, thông qua việc ký kết thỏa thuận với bên xử lý phụ, áp đặt đối với bên '
  'xử lý phụ các nghĩa vụ và yêu cầu tương đương (hoặc cao hơn) so với nghĩa vụ của chính Bên Xử Lý Dữ Liệu trong Thỏa thuận này, '
  'đồng thời giám sát cần thiết và phù hợp đối với bên xử lý phụ.')
NOTE('4.4. (eSIM) Các Bên ghi nhận nhà cung cấp eSIM ở nước ngoài (ví dụ: Consortio) và các đối tác phục vụ vận hành (đơn vị phát '
     'hành hóa đơn điện tử – VNPT; kênh thanh toán – VietinBank/ZaloPay; nền tảng Zalo Mini App) tham gia vào luồng xử lý đơn hàng. '
     'Trong phạm vi các đối tác này xử lý Dữ Liệu Cá Nhân thay mặt Bên Xử Lý Dữ Liệu, họ được coi là bên xử lý phụ và phải tuân thủ '
     'Điều 4 này; việc chuyển dữ liệu cho nhà cung cấp ở nước ngoài còn phải tuân thủ Điều 9 và quy định về chuyển dữ liệu cá nhân '
     'ra nước ngoài. Danh sách bên xử lý phụ được Các Bên thống nhất tại Phụ lục đính kèm.')

# ---------- Article 5 ----------
H('Điều 5. Bảo mật')
P('5.1. Bên Xử Lý Dữ Liệu không được gây ra hoặc cho phép việc tiết lộ hoặc cấp quyền truy cập Dữ Liệu Cá Nhân dưới bất kỳ hình '
  'thức nào cho bất kỳ Nhân sự nào không cần biết Dữ Liệu Cá Nhân cho Mục đích, hoặc cho bất kỳ bên xử lý phụ hay bên thứ ba nào '
  'khác, nếu không có sự chấp thuận trước bằng văn bản của Bên Kiểm Soát Dữ Liệu, trừ khi có yêu cầu khác theo Quy định về Bảo vệ '
  'Dữ liệu cá nhân.')
P('5.2. Trong trường hợp việc tiết lộ hoặc cấp quyền truy cập được Bên Kiểm Soát Dữ Liệu chấp thuận hợp lệ, Bên Xử Lý Dữ Liệu phải '
  'bảo đảm rằng: (a) áp đặt đối với bên xử lý phụ hoặc bên thứ ba các nghĩa vụ và yêu cầu tương đương (hoặc cao hơn) so với Thỏa '
  'thuận này và phù hợp với Quy định về Bảo vệ Dữ liệu cá nhân; và (b) chịu trách nhiệm đối với bất kỳ vi phạm nào do bên xử lý '
  'phụ hoặc bên thứ ba đó gây ra.')

# ---------- Article 6 ----------
H('Điều 6. Nhân sự')
P('6.1. Bên Xử Lý Dữ Liệu phải bảo đảm rằng việc Xử Lý Dữ Liệu Cá Nhân chỉ được thực hiện bởi các Nhân sự chịu trách nhiệm xử lý, '
  'và các Nhân sự đó: (a) đã được đào tạo phù hợp về bảo vệ dữ liệu cá nhân theo quy định, chính sách của Bên Kiểm Soát Dữ Liệu và '
  'Quy định về Bảo vệ Dữ liệu cá nhân; (b) hiểu đầy đủ các nghĩa vụ áp dụng đối với Bên Xử Lý Dữ Liệu và Nhân sự theo Thỏa thuận '
  'này; và (c) đã ký kết các thỏa thuận bảo mật phù hợp, trong đó nghĩa vụ bảo mật vẫn tiếp tục ngay cả sau khi nghỉ việc.')
P('6.2. Bên Xử Lý Dữ Liệu phải yêu cầu Nhân sự tham gia các khóa đào tạo định kỳ về việc Xử Lý Dữ Liệu Cá Nhân và các chương trình '
  'đào tạo nhằm bảo đảm tuân thủ Thỏa thuận này và Quy định về Bảo vệ Dữ liệu cá nhân.')

# ---------- Article 7 ----------
H('Điều 7. Quyền Kiểm Tra')
P('7.1. Bên Kiểm Soát Dữ Liệu có quyền, sau khi thông báo trước cho Bên Xử Lý Dữ Liệu, tiến hành kiểm tra nhằm xác nhận rằng Bên '
  'Xử Lý Dữ Liệu đang thực hiện việc Xử Lý Dữ Liệu Cá Nhân phù hợp với Thỏa thuận này. Việc kiểm tra giới hạn trong phạm vi các hệ '
  'thống vật lý và/hoặc điện tử nơi diễn ra hoạt động Xử Lý Dữ Liệu Cá Nhân và phải tuân thủ các nghĩa vụ bảo mật.')
P('7.2. Bên Xử Lý Dữ Liệu phải, vào bất kỳ thời điểm nào, cung cấp tất cả thông tin cần thiết (ví dụ nội dung các biện pháp bảo mật) '
  'để chứng minh việc tuân thủ Thỏa thuận này theo yêu cầu của Bên Kiểm Soát Dữ Liệu. Nếu phát hiện thiếu sót về bảo mật hoặc không '
  'tuân thủ, Bên Xử Lý Dữ Liệu phải, bằng chi phí của mình, thực hiện mọi biện pháp cần thiết để khắc phục trong thời hạn do Bên Kiểm '
  'Soát Dữ Liệu chỉ định bằng văn bản.')

# ---------- Article 8 ----------
H('Điều 8. Hủy hoặc Xóa Dữ Liệu Cá Nhân')
P('8.1. Bên Xử Lý Dữ Liệu, bằng chi phí của mình, phải vĩnh viễn hủy hoặc xóa Dữ Liệu Cá Nhân đã Xử Lý cho Bên Kiểm Soát Dữ Liệu '
  'và cung cấp văn bản xác nhận đã thực hiện, trước thời hạn do Bên Kiểm Soát Dữ Liệu chỉ định hoặc chậm nhất trong vòng ba (03) '
  'ngày làm việc kể từ khi xảy ra một trong các sự kiện: (a) Hợp Đồng Cơ Sở chấm dứt; (b) việc Xử Lý không còn cần thiết theo Hợp '
  'Đồng Cơ Sở; (c) theo yêu cầu của Bên Kiểm Soát Dữ Liệu; hoặc (d) các trường hợp khác theo Quy định về Bảo vệ Dữ liệu cá nhân.')
NOTE('8.2. (eSIM) Ngoại lệ theo pháp luật chuyên ngành: đối với dữ liệu gắn với hóa đơn, chứng từ kế toán, Bên Xử Lý Dữ Liệu được '
     'lưu giữ trong thời hạn tối thiểu theo quy định của pháp luật về kế toán và hóa đơn (10 năm) ngay cả sau khi Hợp Đồng Cơ Sở '
     'chấm dứt, và chỉ xử lý cho mục đích tuân thủ pháp luật. Cần Pháp chế rà soát để bảo đảm thống nhất với chính sách lưu trữ.')
P('8.3. Việc xóa, hủy phải được thực hiện bằng các biện pháp bảo mật nhằm ngăn chặn việc truy cập trái phép hoặc khôi phục Dữ Liệu '
  'Cá Nhân đã bị xóa, hủy, và phải bảo đảm không xảy ra bất kỳ Vi phạm Dữ liệu nào. Cho đến khi Dữ Liệu Cá Nhân được hủy hoặc xóa '
  'hợp lệ, Bên Xử Lý Dữ Liệu vẫn phải tiếp tục tuân thủ Thỏa thuận này. Nghĩa vụ tại Điều này cũng áp dụng đối với Dữ Liệu Cá Nhân '
  'đã cung cấp cho bất kỳ bên xử lý phụ hoặc bên thứ ba nào.')

# ---------- Article 9 ----------
H('Điều 9. Đánh Giá Tác Động Bảo Vệ Dữ Liệu và Tham Vấn')
P('9.1. Bên Xử Lý Dữ Liệu phải thực hiện đánh giá tác động xử lý dữ liệu cá nhân và đánh giá tác động chuyển dữ liệu cá nhân ra '
  'nước ngoài (nếu có) với tư cách là bên xử lý dữ liệu, đồng thời nộp các hồ sơ liên quan cho Cơ quan có thẩm quyền theo đúng Quy '
  'định về Bảo vệ Dữ liệu cá nhân, trong phạm vi việc Xử Lý Dữ Liệu Cá Nhân do Bên Xử Lý Dữ Liệu thực hiện theo Hợp Đồng Cơ Sở và '
  'Thỏa thuận này.')
P('9.2. Bên Xử Lý Dữ Liệu phải hỗ trợ hợp lý cho Bên Kiểm Soát Dữ Liệu trong việc thực hiện các đánh giá tác động và tham vấn '
  'trước với Cơ quan có thẩm quyền khi Bên Kiểm Soát Dữ Liệu cho rằng cần thiết theo Quy định về Bảo vệ Dữ liệu cá nhân.')
P('9.3. Nếu có vấn đề liên quan đến nội dung của Thỏa thuận này nhưng chưa được quy định hoặc có nghi ngờ trong việc giải thích, '
  'Các Bên sẽ thảo luận thiện chí nhằm đạt được giải pháp phù hợp.')

# ---------- Article 10 ----------
H('Điều 10. Bồi thường')
P('Mỗi Bên (“Bên Bồi Thường”) phải bồi thường và giữ cho Bên còn lại (“Bên Được Bồi Thường”) không bị tổn hại đối với mọi khiếu '
  'nại, yêu cầu, trách nhiệm pháp lý và tổn thất mà Bên Được Bồi Thường phải gánh chịu phát sinh từ hoặc liên quan đến bất kỳ hành '
  'vi vi phạm nghĩa vụ nào theo Thỏa thuận này của Bên Bồi Thường. Các biện pháp khắc phục tại Điều này có tính chất cộng dồn và '
  'không loại trừ, với điều kiện không Bên nào được nhận bồi thường trùng lặp.')

# ---------- Article 11 ----------
H('Điều 11. Chấm dứt')
P('Thỏa thuận này có hiệu lực kể từ Ngày Hiệu Lực và có thể được chấm dứt khi Các Bên thống nhất và Bên Xử Lý Dữ Liệu đã hoàn '
  'thành tất cả các biện pháp quy định tại Điều 8. Việc chấm dứt Hợp Đồng Cơ Sở đồng thời làm phát sinh nghĩa vụ của Bên Xử Lý Dữ '
  'Liệu theo Điều 8.')

# ---------- Article 12 ----------
H('Điều 12. Các quy định khác')
P('12.1. Thông báo. Mọi thông báo, yêu cầu, chấp thuận hoặc liên lạc khác theo Thỏa thuận này phải được lập thành văn bản và gửi '
  'đến địa chỉ/đầu mối của các Bên như sau (hoặc theo thông tin liên hệ mới nhất do bên nhận thông báo):')
P('Bên Kiểm Soát Dữ Liệu (Đại lý): …………………………; Người nhận: …………………; Điện thoại: …………………; Email: …………………', indent=1)
P('Bên Xử Lý Dữ Liệu (GAPIT): Phòng 902, Tầng 9, D10 Giảng Võ, Phường Giảng Võ, Hà Nội; Người nhận: Bà Đỗ Thị Dung; '
  'Điện thoại: 0388871115; Email: dungdt@gapit.com.vn', indent=1)
P('12.2. Luật áp dụng. Thỏa thuận này được điều chỉnh và giải thích theo pháp luật Việt Nam.')
P('12.3. Tính độc lập của các điều khoản. Nếu bất kỳ điều khoản nào bị tuyên bố vô hiệu, không hợp lệ hoặc không thể thi hành, '
  'các điều khoản còn lại vẫn giữ nguyên hiệu lực; điều khoản đó sẽ được sửa đổi để bảo đảm tính hợp lệ trong khi vẫn giữ nguyên '
  'ý định của Các Bên.')
P('12.4. Sửa đổi. Không điều khoản nào của Thỏa thuận này có thể được sửa đổi, thay đổi hoặc chấm dứt trừ khi được lập thành văn '
  'bản và có chữ ký của mỗi Bên.')
P('12.5. Giải quyết tranh chấp. Mọi tranh chấp phát sinh từ hoặc liên quan đến Thỏa thuận này sẽ được giải quyết cuối cùng bằng '
  'trọng tài tại Hà Nội, Việt Nam, tại Trung tâm Trọng tài Quốc tế Việt Nam (VIAC) theo Quy tắc Trọng tài đang có hiệu lực, bởi '
  'một trọng tài viên được chỉ định theo Quy tắc đó. Ngôn ngữ trọng tài là tiếng Việt.')
P('12.6. Mối quan hệ với Hợp Đồng Cơ Sở. Thỏa thuận này là một phần không tách rời của Hợp Đồng Cơ Sở. Trường hợp có mâu thuẫn '
  'về nội dung bảo vệ dữ liệu cá nhân, Thỏa thuận này được ưu tiên áp dụng.')
P('12.7. Bản gốc và hiệu lực. Thỏa thuận này được lập thành 02 (hai) bản gốc có giá trị pháp lý như nhau, mỗi Bên giữ 01 (một) '
  'bản, và có hiệu lực cho đến khi được chấm dứt theo quy định tại đây.')

P('ĐỂ LÀM BẰNG CHỨNG, đại diện có thẩm quyền của Các Bên đã ký kết Thỏa thuận này vào Ngày Hiệu Lực./.', bold=True)

# Signature block
d.add_paragraph()
tbl = d.add_table(rows=2, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
c00 = tbl.cell(0,0).paragraphs[0]; c00.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=c00.add_run('ĐẠI DIỆN BÊN KIỂM SOÁT DỮ LIỆU\n(BÊN B – ĐẠI LÝ)\n(Ký, ghi rõ họ tên, đóng dấu)'); r.bold=True
c01 = tbl.cell(0,1).paragraphs[0]; c01.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=c01.add_run('ĐẠI DIỆN BÊN XỬ LÝ DỮ LIỆU\n(CÔNG TY CỔ PHẦN GAPIT)\n(Ký, ghi rõ họ tên, đóng dấu)'); r.bold=True

# ---------- ANNEX ----------
d.add_page_break()
H('PHỤ LỤC – MÔ TẢ HOẠT ĐỘNG XỬ LÝ DỮ LIỆU', size=13, center=True)
P('Các Bên thống nhất rằng Bên Xử Lý Dữ Liệu sẽ xử lý các loại Dữ Liệu Cá Nhân sau đây của khách hàng cuối (và các nhóm chủ thể '
  'dữ liệu khác nếu có) cho Mục đích đã nêu tại Điều 2.')

H('1. Nhóm chủ thể dữ liệu', size=12)
P('Khách hàng cuối mua Sản phẩm SIM/eSIM du lịch thông qua Bên Kiểm Soát Dữ Liệu (Đại lý); người được Đại lý cử làm đầu mối '
  '(nếu có).')

H('2. Dữ liệu cá nhân cơ bản được xử lý', size=12)
P('Là dữ liệu cá nhân phản ánh các thông tin cá nhân thông thường thường được sử dụng trong các giao dịch, bao gồm (tích √ theo '
  'thực tế thu thập):')
basic = [
 'Họ, tên đệm và tên khai sinh; các tên khác (nếu có);',
 'Số điện thoại;',
 'Địa chỉ liên hệ; địa chỉ thư điện tử (email) – bắt buộc để nhận mã QR/eSIM và hóa đơn;',
 'Số định danh cá nhân; số hộ chiếu (nếu khách hàng cung cấp khi mua/đăng ký);',
 'Thông tin về tài khoản số của cá nhân (ví dụ: tài khoản Zalo);',
 'Mã số thuế cá nhân, thông tin xuất hóa đơn (khi khách hàng yêu cầu xuất hóa đơn);',
 'Các thông tin khác gắn với một cá nhân cụ thể, trừ các thông tin quy định tại Điều 4 Nghị định số 356/2025/NĐ-CP.',
]
for b in basic: P('• ' + b, indent=0.6)

H('3. Dữ liệu liên quan đến giao dịch và hành vi', size=12)
P('Dữ liệu đơn hàng (mã đơn, loại SIM/eSIM, gói cước, số lượng, giá, trạng thái xử lý); thông tin thanh toán/giao dịch (do kênh '
  'thanh toán xử lý); dữ liệu phản ánh hoạt động và lịch sử hoạt động trên không gian mạng phát sinh khi sử dụng nền tảng/dịch vụ.')
NOTE('Lưu ý: trong phạm vi Bên Xử Lý Dữ Liệu nhận/lưu thông tin giao dịch của khách hàng qua tổ chức cung ứng dịch vụ trung gian '
     'thanh toán, dữ liệu này có thể thuộc nhóm dữ liệu cá nhân nhạy cảm theo Luật Bảo vệ Dữ liệu Cá nhân số 91/2025/QH15 và Nghị '
     'định số 356/2025/NĐ-CP; cần áp dụng biện pháp bảo mật nghiêm ngặt và thông báo theo quy định.')

H('4. Dữ liệu cá nhân nhạy cảm', size=12)
P('Về nguyên tắc, dịch vụ SIM/eSIM du lịch không chủ động thu thập dữ liệu cá nhân nhạy cảm. Trường hợp phát sinh (ví dụ: dữ liệu '
  'giao dịch tài chính, dữ liệu vị trí, hình ảnh giấy tờ định danh), Các Bên phải bổ sung mô tả tại Phụ lục này và áp dụng các biện '
  'pháp bảo vệ theo quy định đối với dữ liệu cá nhân nhạy cảm.')

H('5. Mục đích, thời hạn và địa điểm xử lý', size=12)
P('• Mục đích: xử lý đơn hàng, cấp và gửi mã QR/eSIM, xuất hóa đơn, hỗ trợ khách hàng, lưu trữ – tra cứu – đối soát (Điều 2).', indent=0.6)
P('• Thời hạn xử lý/lưu trữ: trong thời hạn Hợp Đồng Cơ Sở và theo chính sách lưu trữ; riêng dữ liệu hóa đơn, chứng từ kế toán lưu '
  'theo thời hạn pháp luật chuyên ngành (tối thiểu 10 năm).', indent=0.6)
P('• Địa điểm: hệ thống CMS và hạ tầng của Bên Xử Lý Dữ Liệu đặt tại Việt Nam; một phần dữ liệu đơn hàng được chuyển cho nhà cung '
  'cấp eSIM ở nước ngoài để cấp eSIM (xem mục 6).', indent=0.6)

H('6. Bên xử lý phụ / Bên thứ ba và chuyển dữ liệu ra nước ngoài', size=12)
tbl2 = d.add_table(rows=1, cols=4); tbl2.style = 'Table Grid'
hdr = tbl2.rows[0].cells
for i,t in enumerate(['Đối tác','Vai trò','Dữ liệu liên quan','Trong/Ngoài nước']):
    rr = hdr[i].paragraphs[0].add_run(t); rr.bold=True
rows = [
 ('Consortio (NCC eSIM)','Cấp/kích hoạt gói eSIM qua API CMS','Thông tin đơn hàng cần thiết để phát hành eSIM','Ngoài nước'),
 ('VNPT','Phát hành hóa đơn điện tử','Tên/đơn vị, MST, địa chỉ, email nhận hóa đơn','Trong nước'),
 ('VietinBank / ZaloPay','Kênh thanh toán/thu hộ','Thông tin giao dịch thanh toán','Trong nước'),
 ('Zalo (VNG)','Nền tảng Zalo Mini App, Zalo OA','Định danh tài khoản Zalo, dữ liệu tương tác','Trong nước'),
]
for r0 in rows:
    cells = tbl2.add_row().cells
    for i,t in enumerate(r0): cells[i].paragraphs[0].add_run(t)
NOTE('Việc chuyển dữ liệu cá nhân cho nhà cung cấp ở nước ngoài (Consortio) phải lập Hồ sơ đánh giá tác động chuyển dữ liệu cá nhân '
     'ra nước ngoài theo Luật Bảo vệ Dữ liệu Cá nhân số 91/2025/QH15 và Nghị định số 356/2025/NĐ-CP, và phải được Bên Kiểm Soát Dữ '
     'Liệu chấp thuận theo Điều 4.')

H('7. Biện pháp bảo vệ dữ liệu (tối thiểu)', size=12)
P('Các biện pháp dưới đây được áp dụng nhất quán với tiêu chuẩn an toàn thông tin mà Bên Xử Lý Dữ Liệu đã thiết lập cho '
  'GAPIT Service Platform / Zalo VAS CMS (zVAS):', italic=True)
for s in [
 'Mã hóa dữ liệu khi truyền tải bằng TLS 1.3 và mã hóa dữ liệu nhạy cảm khi lưu trữ bằng AES-256; tường lửa; phần mềm '
 'phòng chống mã độc;',
 'Phân quyền truy cập theo vai trò (RBAC) và cô lập dữ liệu giữa các khách hàng (multi-tenant isolation) trên hệ thống CMS; '
 'xác thực hai lớp (2FA) bắt buộc đối với tài khoản quản trị (Super_Admin) và hỗ trợ (Support);',
 'Băm mật khẩu bằng bcrypt/argon2; khóa tài khoản sau 05 lần đăng nhập sai; quản lý khóa API, giới hạn tần suất truy cập '
 '(rate limiting) và danh sách IP cho phép (IP whitelist) đối với truy cập của khách hàng doanh nghiệp;',
 'Nhật ký truy cập và hoạt động dạng chỉ-ghi-thêm, không thể chỉnh sửa (immutable audit log) phục vụ truy vết; sao lưu định '
 'kỳ và phương án khôi phục;',
 'Kiểm thử xâm nhập (penetration testing) bắt buộc trước khi đưa lên môi trường vận hành và sau mỗi lần phát hành lớn; '
 'kiểm tra an ninh mạng định kỳ (hàng tháng/quý/năm);',
 'Lưu trữ và xử lý dữ liệu cá nhân trên hạ tầng đặt tại Việt Nam (data residency); quản lý sự đồng ý và bảo đảm quyền yêu '
 'cầu xóa dữ liệu của chủ thể dữ liệu; đào tạo nhân sự định kỳ về bảo vệ dữ liệu cá nhân.',
]:
    P('• ' + s, indent=0.6)

d.save(OUT)
print('SAVED', OUT)
