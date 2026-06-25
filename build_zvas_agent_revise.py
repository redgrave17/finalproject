# -*- coding: utf-8 -*-
"""Sửa HĐ Đại lý zVAS theo góp ý pháp chế (GAPIT=Bên A, không đảo). Highlight phần sửa."""
import copy, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor, Pt
Y=WD_COLOR_INDEX.YELLOW; RED=RGBColor(0xC0,0,0)
SRC='/root/.claude/uploads/0ae05b9e-80f9-5907-8558-033a769d7c2e/7d1fdbb8-HopDong_DaiLy_zVAS_GAPITvDaiLy_1.docx'
OUT='/home/user/finalproject/output/HopDong_DaiLy_zVAS_SUA_THEO_GOPY.docx'
d=docx.Document(SRC); plist=list(d.paragraphs)
def find(needle):
    for p in plist:
        if needle in p.text: return p
    raise ValueError('NOT FOUND: '+needle)
def setp(needle,new):
    p=find(needle)
    if p.runs:
        p.runs[0].text=new
        for r in p.runs[1:]: r.text=''
        r0=p.runs[0]
    else: r0=p.add_run(new)
    r0.font.highlight_color=Y
def insert_after(needle,text,note=False):
    p=find(needle); el=OxmlElement('w:p'); p._p.addnext(el)
    pPr=p._p.find(qn('w:pPr'))
    if pPr is not None:
        npr=copy.deepcopy(pPr)
        for nm in npr.findall(qn('w:numPr')): npr.remove(nm)
        el.append(npr)
    np=Paragraph(el,p._parent)
    if note:
        r=np.add_run('【LƯU Ý】 '); r.bold=True; r.font.color.rgb=RED; r.font.size=Pt(10)
    r2=np.add_run(text); r2.font.highlight_color=Y; r2.font.size=Pt(13); r2.font.name='Times New Roman'

# 1) Lưu ý: văn bản chứng nhận Authorized Reseller Agent của VNG
insert_after('Authorized Reseller Agent',
 'Đề nghị đính kèm/bổ sung văn bản chứng nhận GAPIT là Đại lý phân phối được ủy quyền chính thức (Authorized '
 'Reseller Agent) của VNG đối với các dịch vụ giá trị gia tăng trên nền tảng Zalo, ZingMP3 (Dịch vụ zVAS) làm cơ '
 'sở pháp lý cho tư cách phân phối tại Hợp đồng này.', note=True)

# 2) 4.2 — chiết khấu không được phá giá sàn Zalo (tránh vi phạm quy định VNG)
setp('Chiết khấu từ biên lợi nhuận của chính Đại lý (không công khai)',
 'Ưu đãi từ biên lợi nhuận của chính Đại lý: Bên B có thể dành ưu đãi cho khách hàng từ phần chiết khấu/biên lợi '
 'nhuận của chính mình (mang tính thương lượng riêng, không công khai), với điều kiện trong mọi trường hợp KHÔNG '
 'niêm yết, quảng cáo, chào bán hoặc BÁN Mã code zVAS/Dịch vụ zVAS ở mức giá thấp hơn Giá bán lẻ trực tiếp của Zalo '
 '(Giá sàn công khai) trên bất kỳ kênh nào, nhằm tuân thủ quy định của nhà phát hành (VNG) về việc không niêm yết/'
 'bán Mã code zVAS thấp hơn giá công khai của VNG. Mọi hình thức giảm giá khiến giá bán thực tế thấp hơn Giá sàn '
 'công khai chỉ được thực hiện khi có chấp thuận trước bằng văn bản của Bên A (và của nhà phát hành nếu được yêu cầu).')

# 3) 4.1 & 4.3 — làm rõ, hết mâu thuẫn
setp('cam kết không thay đổi Giá cấp cho Đại lý và mức chiết khấu trong tối thiểu 90',
 'Bên A cam kết không thay đổi Giá cấp cho Đại lý và mức chiết khấu trong tối thiểu 90 (chín mươi) ngày kể từ ngày '
 'áp dụng. Trường hợp điều chỉnh, Bên A thông báo cho Bên B tối thiểu 15 (mười lăm) ngày làm việc trước khi áp dụng '
 'để Hai Bên thống nhất lại giá; giá bán cuối cùng sau điều chỉnh và thỏa thuận được Bên A thông báo và ghi nhận '
 'qua email/CMS.')
setp('Bên A có quyền cập nhật Giá cấp cho Đại lý, Giá bán khuyến nghị và Giá bán lẻ niêm yết',
 'Bên A có quyền cập nhật Giá cấp cho Đại lý, Giá bán khuyến nghị và Giá bán lẻ niêm yết theo từng thời điểm và '
 'thông báo cho Bên B qua email/CMS. Trong thời gian thông báo (tối thiểu 15 ngày làm việc theo Điều 4.1), Hai Bên '
 'có thể thương lượng lại giá; nếu Hai Bên không đạt được thỏa thuận khác, mức giá mới tự động có hiệu lực sau 30 '
 '(ba mươi) ngày kể từ ngày Bên A thông báo.')

# 4) 5.6 — trả trước, kết thúc đơn hàng, không đối soát hàng tháng
setp('Tra cứu và đối chiếu: Bên B chủ động tra cứu lịch sử đơn hàng',
 '5.6. Bản chất trả trước, kết thúc đơn hàng và tra cứu: Hợp đồng vận hành theo cơ chế TRẢ TRƯỚC theo từng đơn; sau '
 'khi Bên B thanh toán đủ và Bên A bàn giao Mã code zVAS theo Điều 5.3, đơn hàng được coi là hoàn tất (kết thúc đơn '
 'hàng). Bên B chủ động tra cứu lịch sử đơn hàng, trạng thái xử lý và hóa đơn trên CMS/zbox.vn; Hai Bên KHÔNG thực '
 'hiện đối soát hàng tháng. Trường hợp có sai lệch, Hai Bên rà soát và điều chỉnh trên tinh thần thiện chí.')

# 5) 5.5 — theo đúng văn bản chị đưa
setp('Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ (thuế suất 10%)',
 '5.5. Hóa đơn GTGT: Bên A xuất hóa đơn GTGT hợp lệ (thuế suất 10%) tương ứng với từng đơn hàng đã thanh toán theo '
 'quy định của pháp luật, trong vòng 03–05 ngày làm việc kể từ ngày tiền thanh toán được ghi có vào tài khoản của '
 'Bên A mở tại Ngân hàng theo thông tin tài khoản quy định trong Hợp đồng.')

# 6) 9.3 & 12.3 — bỏ hoàn tiền/hoàn trả Mã code chưa kích hoạt (vì GAPIT không có quyền này với NCC)
setp('Hoàn tiền Mã code chưa kích hoạt: Trường hợp Bên B không có nhu cầu',
 '9.3. Mã code đã thanh toán không hoàn tiền: Phù hợp với chính sách của nhà phát hành (VNG), Mã code zVAS đã thanh '
 'toán sẽ không được hoàn tiền dưới bất kỳ hình thức nào, kể cả đối với Mã code chưa kích hoạt; ngoại trừ trường hợp '
 'Mã code lỗi thuộc trách nhiệm của Bên A được xử lý theo Điều 9.1. (Lý do: Bên A không được nhà phát hành cho phép '
 'hoàn/đổi Mã code đã thanh toán nên không nhận ràng buộc nghĩa vụ này đối với Bên B.)')
setp('Hoàn trả Mã code chưa kích hoạt: Đối với các Mã code zVAS đã được Bên A bàn giao',
 '12.3. Xử lý Mã code đã bàn giao khi chấm dứt: Đối với các Mã code zVAS đã được Bên A bàn giao và Bên B đã thanh '
 'toán, khi chấm dứt Hợp đồng sẽ không được hoàn trả hoặc hoàn tiền, phù hợp với chính sách không hoàn/đổi của nhà '
 'phát hành (VNG) và Điều 9.3. Bên B chủ động sử dụng hoặc phân phối hết các Mã code đã nhận trước thời điểm chấm dứt.')

# 7) Lưu ý: dùng form trả trước Công ty đã ban hành
insert_after('Hai Bên KHÔNG thực hiện đối soát hàng tháng',
 'Khuyến nghị áp dụng thống nhất các điều khoản trả trước theo mẫu “mua hàng trả trước” do Công ty đã ban hành để '
 'bảo đảm nhất quán giữa các hợp đồng.', note=True)

d.save(OUT); print('SAVED',OUT)
