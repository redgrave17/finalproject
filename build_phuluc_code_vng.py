# -*- coding: utf-8 -*-
"""Phụ lục: Quy trình nhận - xử lý - giao Mã code zVAS cho HĐ Dịch vụ VNG(Bên A)-GAPIT(Bên B).
Không dùng chữ in hoa giữa câu (nhấn mạnh bằng chữ nghiêng)."""
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
OUT='/home/user/finalproject/output/PhuLuc_QuyTrinh_MaCode_zVAS_VNG_GAPIT.docx'
d=docx.Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
NAVY=RGBColor(0x1B,0x2F,0x5A)
def H(t,sz=13,sb=10,center=False,color=NAVY):
    p=d.add_paragraph();r=p.add_run(t);r.bold=True;r.font.size=Pt(sz);r.font.color.rgb=color
    p.paragraph_format.space_before=Pt(sb);p.paragraph_format.space_after=Pt(4)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def P(t,just=True,it=False,ind=None):
    p=d.add_paragraph();r=p.add_run(t);r.italic=it
    if just:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if ind:p.paragraph_format.left_indent=Cm(ind)
    p.paragraph_format.space_after=Pt(3);return p
def shade(cell,hexc):
    cell._tc.get_or_add_tcPr().append(docx.oxml.parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="%s"/>'%hexc))

H('Phụ lục — Quy trình Nhận, Xử lý và Giao Mã code zVAS',sz=14,center=True,sb=0)
P('(Đính kèm và là một phần không tách rời của Hợp đồng Dịch vụ số ……/2026/ZVAS/VNG-…… giữa Công ty Cổ phần Tập '
  'đoàn VNG — “Bên A” và Công ty Cổ phần GAPIT — “Bên B”.)',just=False,it=True)

H('1. Mục đích và phạm vi')
P('Phụ lục này quy định trình tự Bên B (GAPIT) nhận Mã code zVAS từ Bên A (VNG/Zalo), xử lý và quản lý Mã code, và '
  'giao Mã code tới khách hàng của Bên B (đại lý cấp dưới, khách hàng doanh nghiệp và người dùng cuối), bảo đảm '
  'thống nhất trong vận hành và tuân thủ chính sách của Bên A.')

H('2. Sơ đồ tổng quát luồng Mã code')
tbl=d.add_table(rows=1,cols=7); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
flow=[('VNG / Zalo\n(Nhà phát hành – Bên A)','DCE6F7'),('→','FFFFFF'),('GAPIT\n(Nhận & quản lý code – Bên B)','DCE6F7'),
      ('→','FFFFFF'),('Đại lý cấp dưới /\nKhách hàng của GAPIT','DCE6F7'),('→','FFFFFF'),('Người dùng cuối\n(Kích hoạt tại zbox.vn)','DCE6F7')]
for i,(t,c) in enumerate(flow):
    cell=tbl.rows[0].cells[i]; cell.text=''; pr=cell.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=pr.add_run(t); run.bold=(t!='→'); run.font.size=Pt(9 if t!='→' else 14)
    if c!='FFFFFF': shade(cell,c)
P('')

H('3. Bước 1 — Nhận Mã code (Bên B nhận từ Bên A)')
for t in [
 '3.1. Đặt hàng: Bên B gửi yêu cầu đặt hàng cho Bên A qua địa chỉ email đặt hàng được chỉ định tại Hợp đồng; Bên A xác nhận đơn hàng.',
 '3.2. Thanh toán trả trước: Bên B thanh toán 100% giá trị đơn cho Bên A trong vòng 10 (mười) ngày làm việc trước khi sử dụng dịch vụ, vào tài khoản của Bên A theo thông tin quy định tại Hợp đồng. Bên A xuất hóa đơn tài chính hợp lệ sau khi nhận được khoản thanh toán.',
 '3.3. Bàn giao Mã code: Bên A bàn giao Mã code zVAS theo từng đơn (không tách lẻ) trong vòng 03 (ba) ngày làm việc kể từ khi nhận được thanh toán đầy đủ, qua email và/hoặc phương thức khác do Hai Bên thỏa thuận bằng văn bản; đồng thời được cập nhật, tra cứu tại zbox.vn.',
 '3.4. Kiểm tra khi nhận: Bên B kiểm tra số lượng, loại dịch vụ và tình trạng Mã code ngay khi nhận. Mã code bị lỗi, không sử dụng được hoặc đã bị sử dụng trước khi bàn giao được xử lý theo Mục 7 của Phụ lục này.',
]: P(t,ind=0.5)

H('4. Bước 2 — Xử lý và quản lý Mã code')
for t in [
 '4.1. Quản lý code pool: Bên B nhập, lưu trữ và quản lý Mã code trên hệ thống của Bên B, đồng thời tra cứu và đối chiếu tại zbox.vn (mục “Quản lý mã code”): xem danh sách Mã code đã mua gồm Mã code, dịch vụ, ngày tạo, ngày hết hạn, trạng thái (khả dụng, đã kích hoạt, hết hạn); sao chép Mã code; và tải file Excel để đối chiếu.',
 '4.2. Theo dõi tồn kho và hạn kích hoạt: Bên B chủ động theo dõi số lượng Mã code khả dụng và thời hạn hiệu lực kích hoạt của từng Mã code; ưu tiên phân phối các Mã code sắp hết hạn.',
 '4.3. Đối chiếu: trường hợp có sai lệch giữa số liệu của Bên B và hệ thống, Hai Bên rà soát và điều chỉnh trên tinh thần thiện chí.',
]: P(t,ind=0.5)

H('5. Bước 3 — Giao Mã code cho khách hàng (02 mô hình)')
t2=d.add_table(rows=4,cols=3); t2.style='Table Grid'; t2.alignment=WD_TABLE_ALIGNMENT.CENTER
for j,h in enumerate(['Mô hình','Cách giao Mã code','Áp dụng / lưu ý']):
    c=t2.rows[0].cells[j]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); r.font.size=Pt(10); shade(c,'1B2F5A')
rows=[
 ['Quy trình tiêu chuẩn','Bên B giao Mã code cho khách hàng (đại lý cấp dưới hoặc khách hàng doanh nghiệp); khách hàng tự cấp lại Mã code cho người dùng cuối.','Phù hợp khi khách hàng tự quản lý và chăm sóc người dùng cuối của mình.'],
 ['Quy trình nâng cao','Mã code được gửi trực tiếp vào Zalo của người dùng cuối.','Áp dụng theo từng đơn hàng, không tách lẻ.'],
 ['Nghĩa vụ chung','Bên B giao đúng và đủ Mã code cho khách hàng của đơn hàng; bảo mật Mã code, không để lộ, lọt.','Không niêm yết hoặc bán công khai thấp hơn giá bán lẻ công khai của Bên A theo Hợp đồng.'],
]
for i,row in enumerate(rows):
    for j,v in enumerate(row):
        c=t2.rows[i+1].cells[j]; c.text=''; rr=c.paragraphs[0].add_run(v); rr.font.size=Pt(9.5)
        if j==0: rr.bold=True; rr.font.color.rgb=NAVY

H('6. Bước 4 — Kích hoạt (Người dùng cuối)')
for t in [
 '6.1. Người dùng cuối kích hoạt Mã code tại zbox.vn/activate-code: đăng nhập đúng tài khoản Zalo hoặc ZingMP3, nhập Mã code và kích hoạt. Sau khi kích hoạt thành công, tài khoản được sử dụng các tính năng nâng cao tương ứng theo Thời hạn Gói.',
 '6.2. Hỗ trợ: Bên B hỗ trợ, hướng dẫn khách hàng và người dùng cuối trong quá trình kích hoạt; phối hợp với Bên A khi vượt phạm vi xử lý.',
]: P(t,ind=0.5)

H('7. Xử lý Mã code lỗi và hỗ trợ (dẫn chiếu Hợp đồng)')
for t in [
 'Mã code bị lỗi, không sử dụng được hoặc đã bị sử dụng trước thời điểm bàn giao: Bên A hủy hoặc thu hồi và cung cấp lại Mã code khác trong vòng tối đa 24 (hai mươi bốn) giờ làm việc.',
 'Sử dụng Mã code không đúng hướng dẫn dẫn đến tài khoản Zalo/ZingMP3 bị giới hạn hoặc khóa: Bên A không hỗ trợ hủy, thu hồi hoặc hoàn trả Mã code đó.',
 'Mã code đã kích hoạt thành công: không được hoàn tiền dưới bất kỳ hình thức nào, trừ trường hợp lỗi kỹ thuật được Bên A xác nhận.',
 'Giải đáp khiếu nại, thắc mắc của Bên B liên quan đến Mã code trong vòng 24 (hai mươi bốn) giờ làm việc kể từ khi tiếp nhận; nếu không xử lý kịp thời hạn cam kết, Bên A thông báo thời gian xử lý dự kiến.',
]: P('• '+t,ind=0.5)

H('8. Bảo mật và bảo vệ dữ liệu cá nhân')
P('Bên B bảo mật Mã code và các thông tin liên quan theo Hợp đồng. Trường hợp phát sinh việc xử lý dữ liệu cá nhân '
  'của người dùng cuối, các Bên tuân thủ pháp luật về bảo vệ dữ liệu cá nhân (Luật Bảo vệ Dữ liệu Cá nhân số '
  '91/2025/QH15 và Nghị định số 356/2025/NĐ-CP); áp dụng nguyên tắc tối thiểu hóa và khử định danh khi phù hợp. '
  'Dữ liệu định danh phục vụ xác thực (eKYC) là dữ liệu cá nhân nhạy cảm, cần sự đồng ý riêng của chủ thể dữ liệu và '
  'biện pháp bảo mật nghiêm ngặt.')

H('9. Bảng tóm tắt quy trình')
t3=d.add_table(rows=1,cols=4); t3.style='Table Grid'; t3.alignment=WD_TABLE_ALIGNMENT.CENTER
for j,h in enumerate(['Bước','Chủ thể','Hành động chính','Hệ thống / kênh']):
    c=t3.rows[0].cells[j]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); r.font.size=Pt(10); shade(c,'0068FF')
steps=[
 ['1. Đặt & thanh toán','Bên B','Gửi đơn (không tách lẻ) và thanh toán trước 100%','Email đặt hàng'],
 ['2. Nhận code','Bên A → Bên B','Bàn giao Mã code trong 03 ngày làm việc sau thanh toán','Email / zbox.vn'],
 ['3. Quản lý code','Bên B','Tra cứu, tải Excel, theo dõi tồn kho và hạn kích hoạt','zbox.vn — Quản lý mã code'],
 ['4. Giao code','Bên B','Giao cho khách hàng (tiêu chuẩn) hoặc gửi thẳng Zalo người dùng (nâng cao)','Kênh của Bên B / Zalo'],
 ['5. Kích hoạt','Người dùng cuối','Đăng nhập Zalo/ZingMP3, nhập code và kích hoạt','zbox.vn/activate-code'],
]
for row in steps:
    cells=t3.add_row().cells
    for j,v in enumerate(row):
        cells[j].text=''; rr=cells[j].paragraphs[0].add_run(v); rr.font.size=Pt(9.3)
        if j==0: rr.bold=True; rr.font.color.rgb=NAVY

d.add_paragraph()
sg=d.add_table(rows=1,cols=2); sg.alignment=WD_TABLE_ALIGNMENT.CENTER
a=sg.rows[0].cells[0].paragraphs[0]; a.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=a.add_run('Đại diện Bên A\n(Công ty Cổ phần Tập đoàn VNG)\nBà Trần Thị Bảo Vân\n(Ký, ghi rõ họ tên, đóng dấu)'); r.bold=True
b=sg.rows[0].cells[1].paragraphs[0]; b.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=b.add_run('Đại diện Bên B\n(Công ty Cổ phần GAPIT)\nBà Phi Ngọc Cẩm\n(Ký, ghi rõ họ tên, đóng dấu)'); r.bold=True

d.save(OUT); print('SAVED',OUT)
