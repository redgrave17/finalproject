# -*- coding: utf-8 -*-
"""Slide Business Roadmap A-Z cho sản phẩm zVAS:
- Roadmap Now/Next/Later (outcome-driven, tham khảo Adrenalin)
- Biên lợi nhuận GAPIT & Đại lý cấp II
- Order Flow trên zVAS CMS (6 bước, không GapOne)
Branding GAPIT/Zalo lấy từ deck Portfolio."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

C=lambda h:RGBColor.from_string(h)
ZALO=C('0068FF'); NAVY=C('001A40'); BLUE2=C('003F88'); ORANGE=C('F5862C')
TEAL=C('0F6E56'); GREEN=C('22C55E'); PURPLE=C('534AB7'); AMBER=C('BA7517')
CLBLUE=C('185FA5'); GRAY=C('718096'); GRAY2=C('4A5568'); LIGHT=C('F7F8F8')
WHITE=C('FFFFFF'); LN=C('E2E8F0'); INK=C('1A202C')
HEAD='Arial'; BODY='Calibri'
LOGO='/home/user/finalproject/assets/image4.png'

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

def rect(x,y,w,h,fill,line=None,rounded=False,lw=0.75):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                          Inches(x),Inches(y),Inches(w),Inches(h))
    sh.shadow.inherit=False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    return sh

def txt(x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,wrap=True):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.01); tf.margin_bottom=Inches(0.01)
    for i,(runs,sz,sb) in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(0); p.space_before=Pt(sb); p.line_spacing=1.02
        for t,c,b,fn in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c; r.font.name=fn
    return tb

def bullets(x,y,w,h,items,clr,sz=9.5):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.02)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(3); p.space_before=Pt(0); p.line_spacing=1.0
        r=p.add_run(); r.text='▪ '; r.font.size=Pt(sz); r.font.color.rgb=clr; r.font.bold=True; r.font.name=BODY
        r2=p.add_run(); r2.text=it; r2.font.size=Pt(sz); r2.font.color.rgb=GRAY2; r2.font.name=BODY

# ============ TOP BAR ============
rect(0,0,13.333,0.62,LIGHT)
s.shapes.add_picture(LOGO,Inches(0.45),Inches(0.12),height=Inches(0.40))
txt(1.78,0.12,3.0,0.40,[([('×  Zalo VAS',ZALO,True,HEAD)],14,0)],anchor=MSO_ANCHOR.MIDDLE)
rect(10.95,0.10,2.10,0.40,ORANGE,rounded=True)
txt(10.95,0.10,2.10,0.40,[([('BUSINESS ROADMAP A–Z',WHITE,True,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

# ============ TITLE ============
txt(0.45,0.70,9.5,0.46,[([('zVAS — Business Roadmap A → Z',NAVY,True,HEAD)],26,0)])
txt(0.45,1.20,12.6,0.24,[([('Roadmap theo outcome & horizon (Now · Next · Later) · Kinh tế biên lợi nhuận · Order Flow trên zVAS CMS',GRAY,False,BODY)],11,0),
   ([('Khung tham khảo: Adrenalin — “Roadmapping Your Digital Product” (outcome-driven, theme-based, không cố định ngày).',GRAY,False,BODY)],8.5,1)])

# ============ ZONE 1: ROADMAP HORIZONS ============
txt(0.45,1.62,12.0,0.22,[([('1 · ROADMAP SẢN PHẨM — NOW / NEXT / LATER  (theo kết quả/outcome)',GRAY,True,BODY)],10,0)])
cols=[
 (ZALO,'NOW','Q2–Q3 / 2026','Go-live & Validate',
  'Outcome: vận hành B2B ổn định · dòng tiền trả trước · khởi động mạng đại lý',
  ['Hoàn thiện zVAS CMS: catalog, báo giá (prorate + VAT), hóa đơn, kích hoạt, code pool & cảnh báo hạn 6/3/2/1 tháng',
   'Go-live KHDN + ký HĐ Đại lý cấp II; DPA theo PDPA(VNI)',
   'Chuẩn hóa cơ chế trả trước + bảng giá & chiết khấu theo bậc',
   'Đào tạo CSKH cấp 1; quy trình đổi Mã code lỗi (≤ 24h)']),
 (TEAL,'NEXT','Q4/2026 – Q1/2027','Scale & Enable',
  'Outcome: nhân rộng kênh đại lý · tăng MRR · giảm thao tác thủ công',
  ['Mở rộng Đại lý cấp II: RACI, onboarding, KPI doanh số tối thiểu',
   'Portal Đại lý/Affiliate; báo cáo MRR & tồn kho Mã code',
   'Tự động đối soát CSV zbox.vn; auto-select code sắp hết hạn',
   'Thêm bundle (zBusiness+zCloud+zStyle); khuyến mại có kiểm soát giá sàn']),
 (PURPLE,'LATER','2027+','Optimize & Expand',
  'Outcome: tối ưu biên lợi nhuận · mở rộng hệ sinh thái · tích hợp sâu',
  ['API hóa khi VNG mở API → bỏ thao tác CSV thủ công',
   'Cross-sell eSIM / SIM doanh nghiệp; combo đa dịch vụ',
   'Analytics hành vi & dự báo nhu cầu; tối ưu pricing theo bậc',
   'Loyalty cho Đại lý & KHDN; mở rộng vùng/ngành']),
]
cw=4.05; xs=[0.45,4.64,8.83]
for (clr,hz,tf_,tag,oc,items),x in zip(cols,xs):
    rect(x,1.88,cw,0.62,clr,rounded=True)
    txt(x+0.16,1.92,cw-0.3,0.30,[([(hz,WHITE,True,HEAD),('   '+tf_,WHITE,False,BODY)],13,0)])
    txt(x+0.16,2.22,cw-0.3,0.24,[([(tag,WHITE,True,BODY)],10,0)])
    rect(x,2.54,cw,2.18,WHITE,line=LN)
    txt(x+0.12,2.60,cw-0.24,0.40,[([(oc,clr,True,BODY)],8.3,0)])
    bullets(x+0.10,3.00,cw-0.2,1.68,items,clr,sz=9.2)

# ============ ZONE 2: MARGINS (bottom-left) ============
txt(0.45,4.92,7.0,0.22,[([('2 · BIÊN LỢI NHUẬN — GAPIT & ĐẠI LÝ CẤP II',GRAY,True,BODY)],10,0)])
tx=0.45; ty=5.18
colw=[1.55,1.45,1.35,1.55,1.05]; rowh=0.32
hdr=['Bậc SL (code/tháng)','CK Đại lý cấp II','Biên Đại lý II','Biên GAPIT giữ','Tổng biên kênh']
rows=[['50 – 99','5%','5%','15%','20%'],
      ['100 – 299','8%','8%','12%','20%'],
      ['≥ 300','10%','10%','10%','20%']]
# header
cx=tx
for j,htext in enumerate(hdr):
    rect(cx,ty,colw[j],rowh,BLUE2)
    txt(cx,ty,colw[j],rowh,[([(htext,WHITE,True,BODY)],8,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    cx+=colw[j]
# rows
for i,row in enumerate(rows):
    cx=tx; yy=ty+rowh*(i+1)
    for j,val in enumerate(row):
        fill=WHITE if i%2==0 else C('F4F7FB')
        if j==3: fill=C('FFF1E6')   # biên GAPIT cột nổi bật
        rect(cx,yy,colw[j],rowh,fill,line=LN,lw=0.5)
        bold=(j==3); col=ORANGE if j==3 else (NAVY if j==0 else GRAY2)
        txt(cx,yy,colw[j],rowh,[([(val,col,bold,BODY)],8.5,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        cx+=colw[j]
# example + note
txt(0.45,6.30,6.3,0.9,[
 ([('Ví dụ zBusiness 12 tháng — Zalo retail 1.990.000đ · VNG→GAPIT 1.592.000đ (−20%):',NAVY,True,BODY)],8.6,0),
 ([('• Bậc ≥300 (CK 10%): ĐL mua 1.791.000đ → ',GRAY2,False,BODY),('Biên GAPIT 199.000đ (10%)',ORANGE,True,BODY),(' · Biên ĐL 199.000đ (10%)',GREEN,True,BODY)],8.4,1),
 ([('• Bậc 50–99 (CK 5%): ĐL mua 1.890.500đ → ',GRAY2,False,BODY),('Biên GAPIT 298.500đ (15%)',ORANGE,True,BODY),(' · Biên ĐL 99.500đ (5%)',GREEN,True,BODY)],8.4,1),
 ([('Tổng biên kênh ≈20% = chênh giá bán lẻ Zalo vs giá VNG (theo zBusiness 12T; gói khác theo Phụ lục). GAPIT giữ = 20% − CK cấp cho ĐL. Mã code đã mua: không hoàn/đổi (chính sách VNG).',GRAY,False,BODY)],7.6,3),
])

# ============ ZONE 3: ORDER FLOW zVAS CMS (bottom-right) ============
ox=6.95
txt(ox,4.92,6.2,0.22,[([('3 · ORDER FLOW — zVAS CMS (6 bước · không GapOne)',GRAY,True,BODY)],10,0)])
steps=[('01','View Catalog',CLBLUE),('02','Chọn Bundle / Seats',CLBLUE),('03','Báo giá (Prorate + VAT 10%)',PURPLE),
       ('04','Thanh toán + Hóa đơn  ⚠ verify thủ công',AMBER),('05','Subscription Active',GREEN),('06','Pool & Phân bổ seats',GREEN)]
bw=1.93; bh=0.62; gx=0.07; gy=0.30
for k,(n,lbl,clr) in enumerate(steps):
    r=k//3; cI=k%3
    x=ox+cI*(bw+gx); y=5.20+r*(bh+gy)
    rect(x,y,bw,bh,WHITE,line=clr,rounded=True,lw=1.25)
    rect(x,y,0.34,bh,clr,rounded=True)
    txt(x,y,0.34,bh,[([(n,WHITE,True,HEAD)],10,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(x+0.38,y,bw-0.42,bh,[([(lbl,NAVY,True,BODY)],8.2,0)],anchor=MSO_ANCHOR.MIDDLE)
    if cI<2:
        txt(x+bw,y,gx+0.02,bh,[([('→',clr,True,HEAD)],12,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
# down arrow 03->04
txt(ox+2*(bw+gx),5.20+bh,bw,gy,[([('↓',PURPLE,True,HEAD)],12,0)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
# status pipeline
txt(ox,6.40,6.2,0.2,[([('Trạng thái: ',GRAY,True,BODY),
   ('— → draft → draft(quoted) → ',C('4338CA'),False,BODY),('pending_payment',AMBER,True,BODY),
   (' → ',GRAY,False,BODY),('active ✓ → active + pool',C('14532D'),True,BODY)],8.2,0)])
# actor legend (2 hàng x 2 cột)
def dot(x,y,c): rect(x,y,0.12,0.12,c,rounded=True)
legs=[('Client (Owner/Admin)',CLBLUE,ox,6.64),
      ('zVAS CMS — tự động',PURPLE,ox+3.10,6.64),
      ('GAPIT Sales/Admin',TEAL,ox,6.88),
      ('Payment: ZaloPay/Bank · VNPT',AMBER,ox+3.10,6.88)]
for name,c,lx,ly in legs:
    dot(lx,ly+0.03,c)
    txt(lx+0.16,ly-0.02,2.92,0.22,[([(name,GRAY2,False,BODY)],7.8,0)],anchor=MSO_ANCHOR.MIDDLE)

# ============ FOOTER ============
rect(0.45,7.20,12.45,0.22,BLUE2,rounded=True)
txt(0.6,7.20,12.1,0.22,[([('GAPIT JSC  |  gapit.com.vn',WHITE,True,BODY),('     —  Business Roadmap sản phẩm zVAS (Roadmap · Biên lợi nhuận · Order Flow)',WHITE,False,BODY)],8,0)],anchor=MSO_ANCHOR.MIDDLE)

import os
os.makedirs('/home/user/finalproject/output',exist_ok=True)
out='/home/user/finalproject/output/GAPIT_zVAS_Business_Roadmap_AtoZ.pptx'
prs.save(out); print('SAVED',out)
