# -*- coding: utf-8 -*-
"""Cập nhật SOP nội bộ (SOP_QuanLyMaCode_zVAS_NoiBo_2k7_v1) — bổ sung 2 luồng nhận/bàn giao Mã code:
   Luồng 1 = nhận qua CMS (License Pool); Luồng 2 = nhận thủ công qua email (theo Điều 6 HĐ final).
   Toàn bộ nội dung bổ sung được highlight vàng. Xuất bản revised để tải/đưa lên Drive."""
import docx, os
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor, Pt

SRC = '/home/user/finalproject/output/SOP_original.docx'
OUT = '/home/user/finalproject/output/SOP_QuanLyMaCode_zVAS_NoiBo_REVISED_2Luong.docx'
Y = WD_COLOR_INDEX.YELLOW
FN, FS = 'Times New Roman', Pt(13)

d = docx.Document(SRC)

def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

def find_p(needle):
    for p in d.paragraphs:
        if needle in p.text:
            return p
    raise ValueError('NOT FOUND: ' + needle)

def style_run(r, bold=False, hl=True, color=None):
    r.bold = bold
    r.font.name = FN; r.font.size = FS
    if hl: r.font.highlight_color = Y
    if color: r.font.color.rgb = color

def ins_p_before(anchor, segs, bullet=False, space_before=0):
    """segs: list of (text, bold). Chèn 1 đoạn TRƯỚC anchor. Highlight vàng."""
    p = anchor.insert_paragraph_before()
    if space_before:
        p.paragraph_format.space_before = Pt(space_before)
    if bullet:
        try: p.style = d.styles['List Bullet']
        except Exception: pass
    for text, bold in segs:
        style_run(p.add_run(text), bold=bold)
    return p

# Điểm chèn: ngay trước "VI. ĐỒNG BỘ HỆ THỐNG CMS"
anchor = find_p('VI. ĐỒNG BỘ HỆ THỐNG CMS')

# ---------- Heading + intro ----------
ins_p_before(anchor, [('V-A. HAI LUỒNG NHẬN VÀ BÀN GIAO MÃ CODE (bổ sung)', True)], space_before=6)
ins_p_before(anchor, [('Bước 5 (Xuất & bàn giao code) tại Mục V được chi tiết hóa thành hai luồng, tùy theo cách '
    'Đại lý/Khách hàng doanh nghiệp (KHDN) nhận Mã code. Hai luồng dùng chung Bước 3 (xác nhận thanh toán) và '
    'Bước 4 (duyệt xuất); chỉ khác nhau ở khâu xuất, bàn giao, thời điểm chuyển rủi ro và cách lưu bằng chứng.', False)])

# ---------- Luồng 1 ----------
ins_p_before(anchor, [('Luồng 1 — Nhận Mã code qua CMS (License Pool, tự động):', True)], space_before=4)
L1 = [
 'Đối tượng áp dụng: Đại lý/KHDN được GAPIT cấp tài khoản trên CMS/Portal quản lý dịch vụ.',
 'Sau khi Kế toán xác nhận thanh toán (Bước 3) và PM GAPCOM duyệt xuất (Bước 4), Vận hành CMS phân bổ Mã code '
 'vào License Pool của tenant khách hàng; chỉ số allocated_seats tăng tương ứng.',
 'Khách hàng tự truy cập CMS để lấy/kích hoạt Mã code; hệ thống ghi audit log (ai lấy, thời điểm) làm chứng từ.',
 'Bàn giao được coi là hoàn tất khi Mã code được phân bổ vào License Pool và hiển thị trên tài khoản CMS của '
 'khách hàng; rủi ro quản lý Mã code chuyển sang khách hàng kể từ thời điểm này.',
 'Ưu điểm: kiểm soát tồn kho theo thời gian thực, audit log tự động, giảm thao tác thủ công và sai sót.',
]
for t in L1: ins_p_before(anchor, [(t, False)], bullet=True)

# ---------- Luồng 2 ----------
ins_p_before(anchor, [('Luồng 2 — Nhận Mã code thủ công qua email (theo Điều 6 Hợp đồng dịch vụ với KHDN):', True)], space_before=4)
L2 = [
 'Đối tượng áp dụng: Đại lý/KHDN nhận Mã code trực tiếp qua email đã đăng ký, không dùng/không có tài khoản CMS '
 '— đúng theo Điều 6 Hợp đồng.',
 'Sau khi xác nhận thanh toán (Bước 3) và duyệt xuất (Bước 4), Vận hành CMS xuất Mã code và gửi tới đúng địa chỉ '
 'email khách hàng đã cung cấp tại Hợp đồng/Phụ lục 02, trong vòng 03 (ba) ngày làm việc kể từ khi nhận đủ thanh '
 'toán (Điều 6.1).',
 'Khách hàng chịu trách nhiệm bảo mật địa chỉ và hệ thống email dùng để nhận Đơn hàng và Mã code (Điều 6.2).',
 'Việc bàn giao được xác nhận qua xác nhận đơn hàng và xác nhận bàn giao qua email; lập Biên bản nghiệm thu nếu '
 'khách hàng có nhu cầu (Điều 6.3).',
 'Kể từ thời điểm GAPIT gửi Mã code đến đúng email khách hàng đã đăng ký, việc bàn giao được coi là hoàn tất và '
 'mọi rủi ro liên quan đến quản lý, bảo mật, thất lạc Mã code chuyển sang khách hàng (Điều 6.5).',
 'Kiểm soát nội bộ bắt buộc: Vận hành CMS vẫn cập nhật allocated_seats trên License Pool để đối chiếu tồn kho '
 '(Bước 6) dù bàn giao qua email; đồng thời lưu bằng chứng gửi email (log email/biên bản bàn giao) làm chứng từ '
 'đối chiếu.',
]
for t in L2: ins_p_before(anchor, [(t, False)], bullet=True)

# ---------- Bảng so sánh ----------
ins_p_before(anchor, [('Bảng so sánh nhanh hai luồng:', True)], space_before=4)
rows = [
 ('Tiêu chí', 'Luồng 1 — qua CMS', 'Luồng 2 — thủ công qua email'),
 ('Đối tượng áp dụng', 'KH có tài khoản CMS/Portal', 'KH nhận qua email đã đăng ký (Điều 6 HĐ)'),
 ('Cơ chế xuất', 'Phân bổ vào License Pool của tenant', 'Xuất danh sách Mã code, gửi tới email đã đăng ký'),
 ('Thời hạn bàn giao', 'Ngay khi duyệt xuất và phân bổ Pool', '≤ 03 ngày làm việc từ khi nhận đủ tiền (Đ6.1)'),
 ('Thời điểm hoàn tất bàn giao', 'Khi Mã code hiển thị trên tài khoản CMS', 'Khi GAPIT gửi tới đúng email đã đăng ký (Đ6.5)'),
 ('Thời điểm chuyển rủi ro', 'Khi phân bổ vào Pool của khách hàng', 'Khi gửi tới email đã đăng ký (Đ6.5)'),
 ('Bảo mật kênh nhận', 'GAPIT quản trị CMS và phân quyền tài khoản KH', 'KH bảo mật email/tài khoản nhận (Đ6.2)'),
 ('Bằng chứng bàn giao', 'Audit log CMS (ai lấy, thời điểm)', 'Log email + xác nhận đơn/bàn giao qua email (Đ6.3)'),
 ('Đối chiếu tồn kho', 'allocated_seats cập nhật tự động', 'Cập nhật allocated_seats thủ công sau khi gửi'),
 ('Rủi ro cần kiểm soát', 'Cấp/phân quyền tài khoản KH đúng phạm vi', 'Sai/lộ địa chỉ email, thất lạc — bắt buộc lưu log gửi'),
]
tb = d.add_table(rows=len(rows), cols=3)
tb.style = 'Table Grid'
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tb.cell(ri, ci)
        cell.text = ''
        p = cell.paragraphs[0]
        r = p.add_run(val)
        style_run(r, bold=(ri == 0))
        # tô nền header
        if ri == 0:
            shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'D9E2F3'); cell._tc.get_or_add_tcPr().append(shd)
# di chuyển bảng lên trước anchor
anchor._p.addprevious(tb._tbl)

# ---------- ghi chú cập nhật lịch sử ----------
ins_p_before(anchor, [('Ghi chú cập nhật (bản revised): bổ sung Mục V-A tách hai luồng nhận/bàn giao Mã code '
    '(qua CMS và thủ công qua email theo Điều 6 Hợp đồng dịch vụ với KHDN), kèm bảng so sánh — phần highlight '
    'là thông tin bổ sung so với bản gốc.', False)], space_before=4)

os.makedirs('/home/user/finalproject/output', exist_ok=True)
d.save(OUT)
print('SAVED', OUT)
# in lại để kiểm tra
d2 = docx.Document(OUT)
for i, p in enumerate(d2.paragraphs):
    if 'V-A' in p.text or 'Luồng' in p.text:
        print(i, '|', p.text[:80])
print('TABLES', len(d2.tables))
